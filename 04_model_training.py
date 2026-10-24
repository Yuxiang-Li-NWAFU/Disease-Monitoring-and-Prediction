import os
os.environ['PYTHONWARNINGS'] = 'ignore'

import warnings
warnings.filterwarnings('ignore')
warnings.simplefilter('ignore')

import numpy as np
import joblib
import os
import pandas as pd
import json
import collections
import matplotlib
import traceback
import shutil
import itertools

matplotlib.use('Agg')
import matplotlib.pyplot as plt
import seaborn as sns
from datetime import datetime
from sklearn.ensemble import RandomForestClassifier
from sklearn.neural_network import MLPClassifier
from sklearn.tree import DecisionTreeClassifier
from sklearn.svm import SVC
from sklearn.feature_selection import RFECV
from sklearn.model_selection import StratifiedKFold, GridSearchCV, learning_curve
from sklearn.metrics import (accuracy_score, classification_report, confusion_matrix, roc_curve, auc, roc_auc_score,
                             precision_recall_curve, average_precision_score, precision_score, recall_score,
                             f1_score, balanced_accuracy_score)
from imblearn.over_sampling import SMOTE
from sklearn.preprocessing import LabelEncoder, StandardScaler
from sklearn.inspection import permutation_importance
from sklearn.decomposition import PCA
import matplotlib.font_manager as fm
from scipy.stats import chi2, spearmanr
from matplotlib.patches import Ellipse
from mpl_toolkits.mplot3d import Axes3D
import shap
from imblearn.over_sampling import BorderlineSMOTE
from sklearn.feature_selection import mutual_info_classif

import lightgbm as lgb
import xgboost as xgb

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
os.chdir(PROJECT_DIR)

warnings.filterwarnings('ignore', category=UserWarning, module='shap')
warnings.filterwarnings('ignore', category=UserWarning, module='sklearn')
warnings.filterwarnings('ignore', category=UserWarning, module='xgboost')
warnings.filterwarnings('ignore', message='X does not have valid feature names')
warnings.filterwarnings('ignore', message='.*feature names.*')
warnings.filterwarnings('ignore', category=UserWarning, message='.*max_iter.*')
warnings.filterwarnings('ignore', category=UserWarning, message='.*convergence.*')
warnings.filterwarnings('ignore', category=FutureWarning, module='sklearn')
warnings.filterwarnings('ignore', category=FutureWarning, module='pandas')
warnings.filterwarnings('ignore', category=RuntimeWarning, message='.*invalid value.*')
warnings.filterwarnings('ignore', category=RuntimeWarning, message='.*divide by zero.*')
warnings.filterwarnings('ignore', category=UserWarning, module='shap')
warnings.filterwarnings('ignore')
plt.style.use('ggplot')

plt.rcParams['font.family'] = 'Arial'
plt.rcParams['font.size'] = 12
plt.rcParams['axes.labelsize'] = 12
plt.rcParams['axes.titlesize'] = 14
plt.rcParams['legend.fontsize'] = 11
plt.rcParams['xtick.labelsize'] = 10
plt.rcParams['ytick.labelsize'] = 10
plt.rcParams['mathtext.fontset'] = 'stix'
plt.rcParams['mathtext.rm'] = 'Arial'
plt.rcParams['mathtext.it'] = 'Arial:italic'
plt.rcParams['mathtext.bf'] = 'Arial:bold'

def setup_plot_font():
    try:
        fallback_fonts = ['SimHei', 'Microsoft YaHei', 'SimSun', 'Arial Unicode MS']
        for font in fallback_fonts:
            if font in [f.name for f in fm.fontManager.ttflist]:
                plt.rcParams['font.family'] = [font, 'Arial']
                
                return True
        plt.rcParams['font.family'] = 'Arial'
        
        return False
    except Exception as e:
        print(f"Font configuration failed: {e}")
        plt.rcParams['font.family'] = 'Arial'
        return False
plot_font_configured = setup_plot_font()

def convert_types(obj):
    try:
        if isinstance(obj, np.integer):
            return int(obj)
        elif isinstance(obj, np.floating):
            return float(obj)
        elif isinstance(obj, np.bool_):
            return bool(obj)
        elif isinstance(obj, np.str_):
            return str(obj)
        elif isinstance(obj, np.ndarray):
            return obj.tolist()
        elif isinstance(obj, dict):
            return {convert_types(k): convert_types(v) for k, v in obj.items()}
        elif isinstance(obj, list):
            return [convert_types(item) for item in obj]
        elif isinstance(obj, tuple):
            return tuple(convert_types(item) for item in obj)
        elif isinstance(obj, set):
            return [convert_types(item) for item in obj]
        elif isinstance(obj, datetime):
            return obj.isoformat()
        elif hasattr(obj, '__iter__') and not isinstance(obj, str):
            return [convert_types(item) for item in obj]
        else:
            return str(obj)
    except Exception:
        return str(obj)

def plot_confusion_matrix_fixed(y_true, y_pred, classes=None, normalize=False,
                                title="", cmap='Blues', figsize=(10, 8)):
    try:
        y_true = np.array(y_true, dtype=int)
        y_pred = np.array(y_pred, dtype=int)
        cm = confusion_matrix(y_true, y_pred)
        if classes is None:
            unique_classes = np.unique(np.concatenate([y_true, y_pred]))
            classes = [str(cls) for cls in unique_classes]
        else:
            classes = [str(cls) for cls in classes]
        if normalize:
            cm = cm.astype('float') / cm.sum(axis=1)[:, np.newaxis]
            cm = np.nan_to_num(cm, nan=0.0)
            fmt = '.2f'
            vmin, vmax = 0, 1
            title_suffix = " (Normalized)"
        else:
            fmt = 'd'
            vmin, vmax = None, None
            title_suffix = ""
        plt.figure(figsize=figsize)
        sns.heatmap(cm, annot=True, fmt=fmt, cmap=cmap,
                    xticklabels=classes, yticklabels=classes,
                    vmin=vmin, vmax=vmax, cbar_kws={'shrink': 0.8})
        full_title = f"{title}{title_suffix}".strip()
        plt.title(full_title, fontsize=14, fontweight='bold', pad=20, fontfamily='Arial')
        plt.xlabel('Predicted Label', fontsize=12, fontweight='bold', fontfamily='Arial')
        plt.ylabel('True Label', fontsize=12, fontweight='bold', fontfamily='Arial')
        plt.xticks(rotation=45, ha='right', fontfamily='Arial')
        plt.yticks(rotation=0, fontfamily='Arial')
        plt.tight_layout()
        return plt.gcf()
    except Exception as e:
        print(f"Failed to plot confusion matrix: {str(e)}")
        traceback.print_exc()
        return None

def save_confusion_matrix(y_true, y_pred, output_path, classes=None, normalize=False,
                          title="", figsize=(10, 8)):
    fig = plot_confusion_matrix_fixed(y_true, y_pred, classes, normalize, title, figsize=figsize)
    if fig is not None:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        fig.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close(fig)
        
        return True
    return False

def save_confusion_matrix_excel(y_true, y_pred, output_path, classes=None, normalize=False):
    try:
        y_true = np.array(y_true, dtype=int)
        y_pred = np.array(y_pred, dtype=int)
        cm = confusion_matrix(y_true, y_pred)
        if classes is None:
            unique_classes = np.unique(np.concatenate([y_true, y_pred]))
            classes = [str(cls) for cls in unique_classes]
        else:
            classes = [str(cls) for cls in classes]
        if normalize:
            cm = cm.astype('float') / cm.sum(axis=1)[:, np.newaxis]
            cm = np.nan_to_num(cm, nan=0.0)
            sheet_name = "Normalized_Confusion_Matrix"
        else:
            sheet_name = "Confusion_Matrix"
        df = pd.DataFrame(cm, index=classes, columns=classes)
        df.loc['Total'] = df.sum(axis=0)
        df['Total'] = df.sum(axis=1)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name=sheet_name)
        
        return True
    except Exception as e:
        print(f"Failed to save confusion-matrix workbook: {str(e)}")
        traceback.print_exc()
        return False

def remove_highly_correlated_features(X, feature_names, threshold=0.95, priority_features=None):
    df = pd.DataFrame(X, columns=feature_names)
    corr_matrix = df.corr().abs()
    upper = corr_matrix.where(np.triu(np.ones(corr_matrix.shape), k=1).astype(bool))
    to_drop = set()
    for col in upper.columns:
        high_corr_cols = upper.index[upper[col] > threshold].tolist()
        for high_col in high_corr_cols:
            col_priority = priority_features and col in priority_features
            high_priority = priority_features and high_col in priority_features
            if col_priority and not high_priority:
                to_drop.add(high_col)
            elif not col_priority and high_priority:
                to_drop.add(col)
            elif col_priority and high_priority:
                continue
            else:
                to_drop.add(high_col)
    kept_features = [col for col in feature_names if col not in to_drop]
    kept_indices = [feature_names.index(col) for col in kept_features]
    X_reduced = X[:, kept_indices]
    
    
    if len(to_drop) > 0:
        pass
    return X_reduced, kept_features

def plot_full_correlation_heatmap(X, feature_names, priority_features, output_path, threshold=0.95):
    df = pd.DataFrame(X, columns=feature_names)
    corr_matrix = df.corr()
    spore_feats = [f for f in feature_names if f in priority_features]
    other_feats = [f for f in feature_names if f not in priority_features]
    sorted_feats = spore_feats + other_feats
    corr_sorted = corr_matrix.loc[sorted_feats, sorted_feats]
    plt.figure(figsize=(max(12, len(feature_names) * 0.3), max(10, len(feature_names) * 0.3)))
    sns.heatmap(corr_sorted, annot=False, cmap='coolwarm', square=True,
                cbar_kws={'shrink': 0.8}, xticklabels=True, yticklabels=True)
    n_spore = len(spore_feats)
    if n_spore > 0:
        plt.axhline(y=n_spore, color='red', linewidth=2)
        plt.axvline(x=n_spore, color='red', linewidth=2)
        plt.text(n_spore / 2, -2, 'Spore-Related Features', ha='center', fontsize=12, color='darkred')
        plt.text(len(feature_names) + 2, n_spore / 2, 'Other Features', rotation=90, va='center', fontsize=12, color='darkblue')
    plt.title(f'Full Feature Correlation Matrix (threshold > {threshold})', fontsize=14)
    plt.tight_layout()
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    plt.savefig(output_path, dpi=300, bbox_inches='tight')
    plt.close()
    
    high_pairs = []
    for i in range(len(feature_names)):
        for j in range(i + 1, len(feature_names)):
            if abs(corr_matrix.iloc[i, j]) > threshold:
                high_pairs.append((feature_names[i], feature_names[j], corr_matrix.iloc[i, j]))
    if high_pairs:
        
        for a, b, corr in high_pairs[:20]:
            pass
    else:
        pass
    return corr_matrix

def engineer_spore_features(df, spore_col='Spore_Index', interact_cols=None,
                            drop_na=True, log_transform=True):
    df_new = df.copy()
    spore_raw = df_new[spore_col]
    if log_transform:
        df_new[f'{spore_col}_log'] = np.log1p(spore_raw)

        df_new = df_new.drop(columns=[spore_col])
    if interact_cols is not None and len(interact_cols) > 0:
        for col in interact_cols:
            if col in df_new.columns:
                df_new[f'{spore_col}_log_x_{col}'] = df_new[f'{spore_col}_log'] * df_new[col]
    if drop_na:
        df_new = df_new.dropna().reset_index(drop=True)
    else:
        df_new = df_new.fillna(0)
    label_candidates = ['Label', 'target', 'disease_level', 'disease_level_binary']
    label_col = None
    for lc in label_candidates:
        if lc in df_new.columns:
            label_col = lc
            break
    if label_col is not None:
        new_feature_names = [c for c in df_new.columns if c != label_col]
    else:
        new_feature_names = list(df_new.columns)
    return df_new, new_feature_names

def select_top_spore_features(X, y, feature_names, spore_prefix='Spore_Index', top_k=1, method='mutual_info'):
    from sklearn.feature_selection import mutual_info_classif
    from scipy.stats import spearmanr
    spore_features = [f for f in feature_names if spore_prefix in f]
    if not spore_features:
        return X, feature_names
    indices = [feature_names.index(f) for f in spore_features]
    X_spore = X[:, indices]
    if method == 'mutual_info':
        scores = mutual_info_classif(X_spore, y, random_state=42)
    elif method == 'correlation':
        scores = [abs(spearmanr(X_spore[:, i], y).correlation) for i in range(X_spore.shape[1])]
        scores = np.nan_to_num(scores)
    else:
        raise ValueError("method must be 'mutual_info' or 'correlation'")
    sorted_idx = np.argsort(scores)[::-1][:top_k]
    selected_spore = [spore_features[i] for i in sorted_idx]
    other_features = [f for f in feature_names if f not in spore_features]
    final_features = other_features + selected_spore
    final_indices = [feature_names.index(f) for f in final_features]
    X_selected = X[:, final_indices]
    
    return X_selected, final_features

def load_data_and_adjust_labels(file_path, expected_feature_names=None,
                                do_spore_engineering=False, spore_col='Spore_Index',
                                interact_cols=None, is_train=True):
    try:
        df = pd.read_csv(file_path)
        if df.empty:
            raise ValueError("Input data file is empty")
        if do_spore_engineering:
            
            df, new_feature_names = engineer_spore_features(
                df, spore_col=spore_col, interact_cols=interact_cols,
                drop_na=is_train, log_transform=True
            )
            label_col = None
            if 'Label' in df.columns:
                label_col = 'Label'
            elif 'target' in df.columns:
                label_col = 'target'
            else:
                if is_train:
                    label_col = df.columns[-1]
                    print(f" Standard target column not found; using the last column '{label_col}' as the target")
                else:
                    label_col = None
            if label_col is not None:
                labels = df[label_col].values.astype(int)
                df = df.drop(columns=[label_col])
            else:
                labels = None
            if expected_feature_names is not None:
                for col in expected_feature_names:
                    if col not in df.columns:
                        df[col] = 0
                df = df[expected_feature_names]
                feature_names = expected_feature_names
            else:
                feature_names = list(df.columns)
            features = df.values
            if labels is None:
                return features, None, feature_names, None
        else:
            features = df.iloc[:, :-1].values
            labels = df.iloc[:, -1].values.astype(int)
            feature_names = df.columns[:-1].tolist()
            if expected_feature_names is not None:
                common = [c for c in expected_feature_names if c in feature_names]
                indices = [feature_names.index(c) for c in common]
                features = features[:, indices]
                feature_names = common
        le = LabelEncoder()
        labels_encoded = le.fit_transform(labels)
        class_names = le.classes_
        
        return features, labels_encoded, feature_names, class_names
    except Exception as e:
        print(f"Failed to load data: {str(e)}")
        traceback.print_exc()
        return None, None, None, None

def plot_class_distribution(y, title, output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    class_counts = np.bincount(y)
    classes = np.arange(len(class_counts))
    plt.figure(figsize=(10, 6))
    plt.bar(classes, class_counts)
    plt.xlabel('Class', fontfamily='Arial')
    plt.ylabel('Count', fontfamily='Arial')
    plt.title(title, fontfamily='Arial')
    plt.xticks(classes, fontfamily='Arial')
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.savefig(output_path, dpi=300, bbox_inches='tight')
    plt.close()

def calculate_binary_auc(y_true, y_probs, model_classes, fold_num):
    try:
        if y_probs is None or len(y_probs) == 0 or np.any(np.isnan(y_probs)) or np.any(np.isinf(y_probs)):
            return 0.0
        unique_classes = np.unique(y_true)
        if len(unique_classes) <= 1 or y_probs.shape[1] != 2:
            return 0.0
        if len(model_classes) == 2:
            if list(model_classes) == [0, 1]:
                positive_probs = y_probs[:, 1]
            else:
                positive_probs = y_probs[:, 0]
        else:
            return 0.0
        auc_score = roc_auc_score(y_true, positive_probs)
        return auc_score
    except Exception as e:
        print(f"Fold {fold_num}: AUC calculation failed: {str(e)}")
        return 0.0

def plot_roc_curve_binary(y_true, y_probs, model_classes, title, output_path, smooth=False, num_points=200):

    try:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        if len(model_classes) != 2:
            
            return False

        if 1 in list(model_classes):
            positive_index = list(model_classes).index(1)
        else:
            positive_index = 1 if len(model_classes) > 1 else 0
        positive_probs = y_probs[:, positive_index]

        fpr, tpr, thresholds = roc_curve(
            y_true,
            positive_probs,
            drop_intermediate=False
        )
        roc_auc = roc_auc_score(y_true, positive_probs) if len(np.unique(y_true)) > 1 else np.nan

        data_path = output_path.replace('.png', '_data.csv')
        roc_df = pd.DataFrame({
            'False_Positive_Rate': fpr,
            'True_Positive_Rate': tpr,
            'Threshold': thresholds
        })
        roc_df.to_csv(data_path, index=False, encoding='utf-8')
        with open(data_path, 'a', encoding='utf-8') as f:
            f.write(
                f'\n\n# ROC curve metadata:\n'
                f'AUC,{roc_auc}\n'
                f'Positive_Class_Index,{positive_index}\n'
            )

        plt.figure(figsize=(10, 8))
        plt.plot(
            fpr, tpr,
            color='darkorange', lw=2,
            label=f'ROC curve (AUC = {roc_auc:.4f})'
        )
        plt.plot([0, 1], [0, 1], color='navy', lw=2, linestyle='--')
        plt.xlim([0.0, 1.0])
        plt.ylim([0.0, 1.05])
        plt.xlabel('False Positive Rate', fontsize=12, fontfamily='Arial')
        plt.ylabel('True Positive Rate', fontsize=12, fontfamily='Arial')
        plt.title(title, fontsize=14, fontweight='bold', fontfamily='Arial')
        plt.legend(loc="lower right", prop={'family': 'Arial', 'size': 12})
        plt.grid(True, alpha=0.3)
        plt.tight_layout()
        plt.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close()
        
        return True
    except Exception as e:
        print(f"Failed to plot ROC curve: {str(e)}")
        traceback.print_exc()
        return False

def plot_pr_curve_binary(y_true, y_probs, model_classes, title, output_path, smooth=True, num_points=200):
    try:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        if len(model_classes) != 2:
            
            return False
        if list(model_classes) == [0, 1]:
            positive_probs = y_probs[:, 1]
        else:
            positive_probs = y_probs[:, 0]

        precision, recall, thresholds = precision_recall_curve(y_true, positive_probs)
        avg_precision = average_precision_score(y_true, positive_probs)

        if len(thresholds) > 0:
            pr_data = pd.DataFrame({
                'Precision': precision[:-1],
                'Recall': recall[:-1],
                'Threshold': thresholds
            })
        else:
            pr_data = pd.DataFrame({'Precision': precision, 'Recall': recall})
        data_path = output_path.replace('.png', '_data.csv')
        pr_data.to_csv(data_path, index=False, encoding='utf-8')
        with open(data_path, 'a', encoding='utf-8') as f:
            f.write(f'\n\n# PR curve metadata:\nAverage_Precision,{avg_precision}\nPositive_Class_Index,{1 if 1 in model_classes else 0}\n')

        plt.figure(figsize=(10, 8))
        if smooth and len(recall) > 2:
            from scipy.interpolate import interp1d
            sorted_idx = np.argsort(recall)
            recall_sorted = recall[sorted_idx]
            precision_sorted = precision[sorted_idx]
            unique_recall = []
            unique_precision = []
            for r, p in zip(recall_sorted, precision_sorted):
                if not unique_recall or r != unique_recall[-1]:
                    unique_recall.append(r)
                    unique_precision.append(p)
                else:
                    unique_precision[-1] = p
            n_unique = len(unique_recall)
            if n_unique >= 4:
                f_interp = interp1d(unique_recall, unique_precision, kind='cubic', fill_value='extrapolate')
                recall_smooth = np.linspace(0, 1, num_points)
                precision_smooth = f_interp(recall_smooth)
                precision_smooth = np.clip(precision_smooth, 0, 1)
                plt.plot(recall_smooth, precision_smooth, color='darkgreen', lw=2, label=f'PR curve (AP = {avg_precision:.4f})')
            elif n_unique >= 2:
                f_interp = interp1d(unique_recall, unique_precision, kind='linear', fill_value='extrapolate')
                recall_smooth = np.linspace(0, 1, num_points)
                precision_smooth = f_interp(recall_smooth)
                precision_smooth = np.clip(precision_smooth, 0, 1)
                plt.plot(recall_smooth, precision_smooth, color='darkgreen', lw=2, label=f'PR curve (AP = {avg_precision:.4f})')
            else:
                plt.plot(recall, precision, color='darkgreen', lw=2, label=f'PR curve (AP = {avg_precision:.4f})')
        else:
            plt.plot(recall, precision, color='darkgreen', lw=2, label=f'PR curve (AP = {avg_precision:.4f})')

        plt.xlabel('Recall', fontsize=12, fontfamily='Arial')
        plt.ylabel('Precision', fontsize=12, fontfamily='Arial')
        plt.title(title, fontsize=14, fontweight='bold', fontfamily='Arial')
        plt.legend(loc="lower left", prop={'family': 'Arial', 'size': 12})
        plt.grid(True, alpha=0.3)
        plt.xlim([0.0, 1.0])
        plt.ylim([0.0, 1.05])
        plt.tight_layout()
        plt.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close()
        
        return True
    except Exception as e:
        print(f"Failed to plot PR curve: {str(e)}")
        traceback.print_exc()
        return False

def compute_binary_metrics(y_true, y_pred, y_prob):

    cm = confusion_matrix(y_true, y_pred, labels=[0, 1])
    tn, fp, fn, tp = cm.ravel()

    accuracy = accuracy_score(y_true, y_pred)
    precision = precision_score(y_true, y_pred, zero_division=0)
    recall = recall_score(y_true, y_pred, zero_division=0)
    f1 = f1_score(y_true, y_pred, zero_division=0)
    ba = balanced_accuracy_score(y_true, y_pred)

    specificity = tn / (tn + fp) if (tn + fp) > 0 else 0.0
    sensitivity = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    gmean = float(np.sqrt(sensitivity * specificity))

    if len(np.unique(y_true)) > 1:
        roc_auc = roc_auc_score(y_true, y_prob)
    else:
        roc_auc = np.nan

    return {
        'TN': int(tn),
        'FP': int(fp),
        'FN': int(fn),
        'TP': int(tp),
        'Accuracy': float(accuracy),
        'Precision': float(precision),
        'Recall': float(recall),
        'F1-score': float(f1),
        'BA': float(ba),
        'G-mean': float(gmean),
        'ROC-AUC': float(roc_auc) if np.isfinite(roc_auc) else None
    }

def analyze_feature_selection_consistency(all_results, result_dir, method='rf'):
    
    all_selected = []
    for res in all_results:
        if method == 'rf' and 'rf_selected_features' in res:
            all_selected.extend(res['rf_selected_features'])
        elif method == 'rfe' and 'rfe_selected_features' in res:
            all_selected.extend(res['rfe_selected_features'])
        elif 'selected_features' in res:
            all_selected.extend(res['selected_features'])
    if not all_selected:
        return []
    counter = collections.Counter(all_selected)
    total = len(all_results)
    df = pd.DataFrame({
        'feature': list(counter.keys()),
        'count': list(counter.values()),
        'frequency': [c / total for c in counter.values()]
    }).sort_values('frequency', ascending=False)
    core = df[df['frequency'] > 0.5]['feature'].tolist()
    os.makedirs(os.path.join(result_dir, 'core_features'), exist_ok=True)
    df.to_csv(os.path.join(result_dir, 'core_features/feature_selection_frequency.csv'), index=False)
    with open(os.path.join(result_dir, 'core_features/core_features.json'), 'w') as f:
        json.dump({'core_features': core, 'total_folds': total, 'threshold': 0.5}, f, indent=4)
    plt.figure(figsize=(12, 8))
    top = df.head(20)
    plt.barh(top['feature'], top['frequency'])
    plt.xlabel('Selection Frequency', fontfamily='Arial')
    plt.title(f'Top 20 Features by Selection Frequency ({method.upper()})', fontfamily='Arial')
    plt.gca().invert_yaxis()
    plt.tight_layout()
    plt.savefig(os.path.join(result_dir, 'core_features/feature_selection_frequency.png'), dpi=300, bbox_inches='tight')
    plt.close()
    
    return core

def generate_final_report(result_dir, config, n_total_features, n_core_features,
                          cv_accuracy, train_accuracy, test_accuracy, test_auc,
                          model_type='RF'):
    
    os.makedirs(os.path.join(result_dir, "summary"), exist_ok=True)
    report = {
        "project_name": f"{config['feature_selector'].upper()}+PCA+{config['classifier'].upper()} model training report",
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "training_config": config,
        "data_statistics": {
            "total_features": n_total_features,
            "core_features": n_core_features,
            "feature_reduction_ratio": f"{(n_total_features - n_core_features) / n_total_features * 100:.1f}%" if n_core_features > 0 else "N/A",
            "feature_retention_ratio": f"{n_core_features / n_total_features * 100:.1f}%" if n_core_features > 0 else "N/A"
        },
        "model_performance": {
            "mean_cross_validation_accuracy": f"{cv_accuracy:.4f}" if cv_accuracy else "N/A",
            "training_accuracy": f"{train_accuracy:.4f}" if train_accuracy else "N/A",
            "test_accuracy": f"{test_accuracy:.4f}" if test_accuracy else "N/A",
            "test_auc": f"{test_auc:.4f}" if test_auc else "N/A"
        },
        "model_information": {
            "model_type": f"{config['feature_selector'].upper()} feature selection + PCA + {config['classifier'].upper()}",
            "feature_selection_method": "Random forest cumulative importance" if config['feature_selector'] == 'rf' else "Recursive feature elimination (RFECV)",
            "dimensionality_reduction_method": "PCA",
            "classifier": config['classifier'].upper()
        }
    }
    if train_accuracy and test_accuracy:
        gap = train_accuracy - test_accuracy
        report["model_performance"]["generalization_gap"] = f"{gap:.4f}"
        report["model_performance"]["generalization_assessment"] = "Good" if abs(gap) < 0.1 else ("Potential overfitting" if gap > 0.1 else "Potential underfitting")
    with open(os.path.join(result_dir, "summary", "final_report.json"), 'w', encoding='utf-8') as f:
        json.dump(convert_types(report), f, indent=4, ensure_ascii=False)
    
    return report

def create_performance_summary_plot(fold_accuracies, fold_auc_scores, result_dir):
    
    try:
        os.makedirs(os.path.join(result_dir, "summary"), exist_ok=True)
        if fold_auc_scores and len(fold_auc_scores) > 0:
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))
        else:
            fig, ax1 = plt.subplots(1, 1, figsize=(12, 5))
            ax2 = None
        if fold_accuracies and len(fold_accuracies) > 0:
            mean_acc = np.mean(fold_accuracies)
            std_acc = np.std(fold_accuracies)
            ax1.bar(range(1, len(fold_accuracies) + 1), fold_accuracies, alpha=0.7, color='skyblue')
            ax1.axhline(mean_acc, color='red', linestyle='--', linewidth=2,
                        label=f'Mean: {mean_acc:.4f} ± {std_acc:.4f}')
            ax1.set_xlabel('Fold', fontfamily='Arial')
            ax1.set_ylabel('Accuracy', fontfamily='Arial')
            ax1.set_title('Cross-Validation Accuracy by Fold', fontfamily='Arial')
            ax1.legend(prop={'family': 'Arial', 'size': 10})
            ax1.grid(True, alpha=0.3)
        if ax2 is not None and fold_auc_scores and len(fold_auc_scores) > 0:
            mean_auc = np.mean(fold_auc_scores)
            std_auc = np.std(fold_auc_scores)
            ax2.bar(range(1, len(fold_auc_scores) + 1), fold_auc_scores, alpha=0.7, color='lightgreen')
            ax2.axhline(mean_auc, color='red', linestyle='--', linewidth=2,
                        label=f'Mean: {mean_auc:.4f} ± {std_auc:.4f}')
            ax2.set_xlabel('Fold', fontfamily='Arial')
            ax2.set_ylabel('AUC', fontfamily='Arial')
            ax2.set_title('Cross-Validation AUC by Fold', fontfamily='Arial')
            ax2.legend(prop={'family': 'Arial', 'size': 10})
            ax2.grid(True, alpha=0.3)
        plt.tight_layout()
        plt.savefig(os.path.join(result_dir, "summary/performance_summary.png"), dpi=300, bbox_inches='tight')
        plt.close()
        
    except Exception as e:
        print(f"Failed to create the performance summary plot: {str(e)}")

def save_training_config(config, result_dir):
    
    try:
        os.makedirs(os.path.join(result_dir, "config"), exist_ok=True)
        config_path = os.path.join(result_dir, "config/training_config.json")
        with open(config_path, 'w', encoding='utf-8') as f:
            json.dump(convert_types(config), f, indent=4, ensure_ascii=False)
        
    except Exception as e:
        print(f"Failed to save training configuration: {str(e)}")

def save_smote_distribution_excel(y_before, y_after, output_path, title="SMOTE Class Distribution"):
    before_counts = np.bincount(y_before)
    after_counts = np.bincount(y_after)
    max_class = max(len(before_counts), len(after_counts))
    before_counts = np.pad(before_counts, (0, max_class - len(before_counts)), 'constant')
    after_counts = np.pad(after_counts, (0, max_class - len(after_counts)), 'constant')
    diff = after_counts - before_counts
    diff_percent = (diff / before_counts * 100).round(2)
    diff_percent = np.where(np.isinf(diff_percent), 0, diff_percent)
    df = pd.DataFrame({
        'Class': list(range(max_class)),
        'Before_SMOTE': before_counts,
        'After_SMOTE': after_counts,
        'Change': diff,
        'Change_%': diff_percent
    })
    df.to_excel(output_path, index=False, sheet_name='SMOTE_Change')
    

class PCAAnalyzer:
    def __init__(self, n_components=None, variance_threshold=0.95, random_state=42):
        self.n_components = n_components
        self.variance_threshold = variance_threshold
        self.random_state = random_state
        self.pca = None
        self.scaler = None
        self.original_feature_names = None
        self.X_pca = None
        self.labels = None

    def _format_class_label(self, label):
        if isinstance(label, (int, float, np.integer)):
            return str(int(label))
        elif isinstance(label, str) and label.startswith('G'):
            return label[1:] if label[1:].isdigit() else label
        else:
            return str(label)

    def fit_transform(self, X, feature_names=None, labels=None):
        self.original_feature_names = feature_names
        self.labels = labels
        self.scaler = StandardScaler()
        X_scaled = self.scaler.fit_transform(X)
        if self.n_components is None:
            temp_pca = PCA(random_state=self.random_state)
            temp_pca.fit(X_scaled)
            cumulative_variance = np.cumsum(temp_pca.explained_variance_ratio_)
            self.n_components = np.argmax(cumulative_variance >= self.variance_threshold) + 1
            self.n_components = max(3, min(self.n_components, X.shape[1]))
            
            
            
        self.pca = PCA(n_components=self.n_components, random_state=self.random_state)
        self.X_pca = self.pca.fit_transform(X_scaled)
        return self.X_pca

    def transform(self, X):
        if self.pca is None or self.scaler is None:
            raise ValueError("PCA has not been fitted. Call fit_transform first.")
        X_scaled = self.scaler.transform(X)
        return self.pca.transform(X_scaled)

    def analyze_feature_importance(self):
        if self.pca is None:
            raise ValueError("PCA has not been fitted")
        loadings = self.pca.components_.T
        explained_variance = self.pca.explained_variance_ratio_
        weighted = np.sum(np.abs(loadings) * explained_variance, axis=1)
        max_imp = np.max(np.abs(loadings) * explained_variance, axis=1)
        vip = np.sum(loadings ** 2 * explained_variance, axis=1)
        weighted = weighted / np.max(weighted)
        max_imp = max_imp / np.max(max_imp)
        vip = vip / np.max(vip)
        composite = (weighted * 0.33 + max_imp * 0.33 + vip * 0.33)
        return {
            'weighted': weighted,
            'max': max_imp,
            'vip': vip,
            'composite': composite,
            'loadings': loadings,
            'explained_variance': explained_variance,
            'pc_names': [f'PC{i + 1}' for i in range(self.n_components)]
        }

    def plot_pca_with_confidence_ellipses(self, figsize=(12, 10), confidence_level=0.95, save_path=None, dpi=300):
        if self.X_pca is None or self.labels is None:
            print(" PCA data or label information is missing")
            return None, None
        try:
            fig, ax = plt.subplots(figsize=figsize)
            unique_labels = np.unique(self.labels)
            colors = plt.cm.tab10(np.linspace(0, 1, len(unique_labels)))
            scale = np.sqrt(chi2.ppf(confidence_level, 2))
            ellipse_count = 0
            for i, label in enumerate(unique_labels):
                mask = self.labels == label
                group_data = self.X_pca[mask]
                if len(group_data) < 2:
                    continue
                formatted_label = self._format_class_label(label)
                ax.scatter(group_data[:, 0], group_data[:, 1],
                           color=colors[i], label=f'Class {formatted_label}',
                           s=60, alpha=0.7, edgecolors='black', linewidth=0.5)
                if len(group_data) >= 3:
                    try:
                        mean = np.mean(group_data[:, :2], axis=0)
                        cov = np.cov(group_data[:, 0], group_data[:, 1])
                        if np.linalg.det(cov) > 1e-10:
                            eigenvalues, eigenvectors = np.linalg.eig(cov)
                            if np.all(eigenvalues > 1e-10):
                                angle = np.degrees(np.arctan2(eigenvectors[1, 0], eigenvectors[0, 0]))
                                width = 2 * np.sqrt(eigenvalues[0]) * scale
                                height = 2 * np.sqrt(eigenvalues[1]) * scale
                                ellipse = Ellipse(xy=mean, width=width, height=height,
                                                  angle=angle, edgecolor=colors[i],
                                                  facecolor='none', linestyle='--', linewidth=2, alpha=0.7)
                                ax.add_patch(ellipse)
                                ellipse_count += 1
                    except Exception:
                        continue
            pc1_var = self.pca.explained_variance_ratio_[0] * 100 if self.pca is not None else 0
            pc2_var = self.pca.explained_variance_ratio_[1] * 100 if self.pca is not None else 0
            ax.set_xlabel(f'PC1 ({pc1_var:.2f}%)', fontsize=12, fontfamily='Arial')
            ax.set_ylabel(f'PC2 ({pc2_var:.2f}%)', fontsize=12, fontfamily='Arial')
            ax.set_title(f'PCA with {confidence_level * 100:.0f}% Confidence Ellipses', fontsize=14, fontweight='bold',
                         fontfamily='Arial')
            ax.grid(True, alpha=0.3)
            legend = ax.legend()
            for text in legend.get_texts():
                text.set_fontproperties(fm.FontProperties(family='Arial', size=11))
            plt.tight_layout()
            if save_path:
                plt.savefig(save_path, dpi=dpi, bbox_inches='tight')
                
            return fig, ax
        except Exception as e:
            print(f"Failed to generate the PCA plot with confidence ellipses: {str(e)}")
            return None, None

    def plot_cumulative_variance(self, threshold=0.95, figsize=(12, 6), save_path=None, dpi=300):
        if self.pca is None:
            
            return None, None, None

        try:
            explained_variances = self.pca.explained_variance_ratio_
            cumulative_variance = np.cumsum(explained_variances)

            n_components_threshold = np.argmax(cumulative_variance >= threshold) + 1

            components_to_show = min(10, len(explained_variances))
            x = np.arange(1, components_to_show + 1)

            fig, ax = plt.subplots(figsize=figsize)

            bars = ax.bar(
                x,
                explained_variances[:components_to_show] * 100,
                color='steelblue',
                alpha=0.7,
                label='Individual contribution'
            )

            for i, v in enumerate(explained_variances[:components_to_show]):
                ax.text(
                    i + 1,
                    v * 100 + 0.5,
                    f'{v * 100:.1f}%',
                    ha='center',
                    va='bottom',
                    fontsize=9,
                    fontfamily='Arial'
                )

            cumulative_line, = ax.plot(
                x,
                cumulative_variance[:components_to_show] * 100,
                'r-o',
                linewidth=2,
                markersize=6,
                label='Cumulative'
            )

            threshold_line = ax.axhline(
                y=threshold * 100,
                color='red',
                linestyle='--',
                linewidth=1.8,
                alpha=0.7,
                label=f'{threshold * 100:.0f}% threshold'
            )

            selected_x = n_components_threshold
            selected_y = cumulative_variance[n_components_threshold - 1] * 100

            if selected_x <= components_to_show:
                ax.plot(
                    selected_x,
                    selected_y,
                    'o',
                    color='red',
                    markersize=10,
                    zorder=5
                )

                ax.text(
                    selected_x + 0.20,
                    selected_y - 3,
                    f'Components: {n_components_threshold}\n'
                    f'Variance: {selected_y:.1f}%',
                    fontsize=10,
                    fontfamily='Arial',
                    bbox=dict(
                        facecolor='white',
                        alpha=0.85,
                        edgecolor='lightgray'
                    )
                )

            ax.set_xlabel(
                'Principal Component',
                fontsize=12,
                fontfamily='Arial'
            )

            ax.set_ylabel(
                'Explained Variance (%)',
                fontsize=12,
                fontfamily='Arial'
            )

            ax.set_title(
                f'Top {components_to_show} Principal Components',
                fontsize=14,
                fontweight='bold',
                fontfamily='Arial'
            )

            ax.set_xticks(x)

            ax.set_ylim(
                0,
                max(
                    100,
                    np.max(cumulative_variance[:components_to_show] * 100) + 5
                )
            )

            ax.grid(
                True,
                alpha=0.3,
                axis='y'
            )

            ax.legend(
                handles=[
                    cumulative_line,
                    bars,
                    threshold_line
                ],
                labels=[
                    'Cumulative',
                    'Individual contribution',
                    f'{threshold * 100:.0f}% threshold'
                ],
                loc='upper left',
                prop={
                    'family': 'Arial',
                    'size': 10
                }
            )

            plt.tight_layout()

            if save_path:
                plt.savefig(
                    save_path,
                    dpi=dpi,
                    bbox_inches='tight'
                )
                

            return fig, ax, n_components_threshold

        except Exception as e:
            print(f"Failed to plot cumulative variance: {str(e)}")
            return None, None, None

    def plot_3d_pca(self, figsize=(12, 10), save_path=None, dpi=300):
        if self.X_pca is None or self.X_pca.shape[1] < 3:
            print("Cannot create a 3D plot: insufficient data dimensions")
            return None, None
        if self.labels is None:
            self.labels = np.zeros(self.X_pca.shape[0])
        try:
            fig = plt.figure(figsize=figsize)
            ax = fig.add_subplot(111, projection='3d')
            unique_labels = np.unique(self.labels)
            colors = plt.cm.tab20(np.linspace(0, 1, len(unique_labels)))
            for i, label in enumerate(unique_labels):
                mask = self.labels == label
                group_data = self.X_pca[mask]
                formatted_label = self._format_class_label(label)
                ax.scatter(group_data[:, 0], group_data[:, 1], group_data[:, 2],
                           c=[colors[i]], label=f'Class {formatted_label}',
                           alpha=0.7, s=50, edgecolors='w', depthshade=True)
            pc1_var = self.pca.explained_variance_ratio_[0] * 100 if self.pca is not None and len(
                self.pca.explained_variance_ratio_) >= 3 else 0
            pc2_var = self.pca.explained_variance_ratio_[1] * 100 if self.pca is not None and len(
                self.pca.explained_variance_ratio_) >= 3 else 0
            pc3_var = self.pca.explained_variance_ratio_[2] * 100 if self.pca is not None and len(
                self.pca.explained_variance_ratio_) >= 3 else 0
            ax.set_xlabel(f'PC1 ({pc1_var:.2f}%)', fontsize=12, labelpad=10, fontfamily='Arial')
            ax.set_ylabel(f'PC2 ({pc2_var:.2f}%)', fontsize=12, labelpad=10, fontfamily='Arial')
            ax.set_zlabel(f'PC3 ({pc3_var:.2f}%)', fontsize=12, labelpad=10, fontfamily='Arial')
            ax.set_title('3D PCA Scatter Plot', fontsize=14, fontweight='bold', pad=20, fontfamily='Arial')
            legend = ax.legend(title='Class', bbox_to_anchor=(1.05, 1), loc='upper left')
            for text in legend.get_texts():
                text.set_fontproperties(fm.FontProperties(family='Arial', size=11))
            legend.get_title().set_fontproperties(fm.FontProperties(family='Arial', size=12))
            ax.view_init(elev=20, azim=45)
            plt.tight_layout()
            if save_path:
                plt.savefig(save_path, dpi=dpi, bbox_inches='tight')
                
            return fig, ax
        except Exception as e:
            print(f"Failed to plot the 3D PCA scatter plot: {str(e)}")
            return None, None

    def save_analysis(self, result_dir, prefix=""):
        os.makedirs(os.path.join(result_dir, "pca_analysis"), exist_ok=True)
        try:
            ellipse_path = os.path.join(result_dir, f"pca_analysis/{prefix}pca_with_confidence_ellipses.png")
            fig1, _ = self.plot_pca_with_confidence_ellipses(save_path=ellipse_path, dpi=300)
            if fig1: plt.close(fig1)
            cum_path = os.path.join(result_dir, f"pca_analysis/{prefix}pca_cumulative_variance.png")
            fig2, _, _ = self.plot_cumulative_variance(save_path=cum_path, dpi=300)
            if fig2: plt.close(fig2)
            if self.X_pca is not None and self.X_pca.shape[1] >= 3:
                pca_3d_path = os.path.join(result_dir, f"pca_analysis/{prefix}pca_3d_scatter.png")
                fig3, _ = self.plot_3d_pca(save_path=pca_3d_path, dpi=300)
                if fig3: plt.close(fig3)
            if hasattr(self, 'original_feature_names') and self.original_feature_names is not None:
                loadings_df = pd.DataFrame(
                    self.pca.components_.T,
                    columns=[f'PC{i + 1}' for i in range(self.n_components)],
                    index=self.original_feature_names
                )
                loadings_df.to_csv(os.path.join(result_dir, f"pca_analysis/{prefix}pca_loadings.csv"))
            
            return self.analyze_feature_importance()
        except Exception as e:
            print(f"Failed to save PCA analysis results: {str(e)}")
            return None

class RFSelector:

    def __init__(self, n_estimators=100, random_state=42, importance_threshold='median'):
        self.n_estimators = n_estimators
        self.random_state = random_state
        self.importance_threshold = importance_threshold
        self.rf = None
        self.selected_indices = None
        self.selected_features = None
        self.feature_importances = None

    def fit(self, X, y, feature_names=None):
        self.rf = RandomForestClassifier(
            n_estimators=self.n_estimators,
            random_state=self.random_state,
            n_jobs=-1,
            class_weight='balanced'
        )
        self.rf.fit(X, y)

        self.feature_importances = np.asarray(self.rf.feature_importances_, dtype=float)
        indices = np.argsort(self.feature_importances)[::-1]
        cumulative = np.cumsum(self.feature_importances[indices])

        n_selected = np.argmax(cumulative >= 0.85) + 1
        n_selected = max(10, n_selected)
        n_selected = min(n_selected, len(self.feature_importances))
        self.selected_indices = indices[:n_selected]

        if feature_names is not None:
            self.selected_features = [feature_names[i] for i in self.selected_indices]
        else:
            self.selected_features = self.selected_indices
        return self

    def transform(self, X):
        if self.selected_indices is None:
            raise ValueError("Feature selector has not been fitted")
        return X[:, self.selected_indices]

    def get_feature_importance_df(self, feature_names=None):

        if self.feature_importances is None:
            raise ValueError("Feature importance has not been computed")
        if feature_names is None:
            feature_names = [f'Feature_{i}' for i in range(len(self.feature_importances))]
        if len(feature_names) != len(self.feature_importances):
            raise ValueError("The number of feature_names does not match the number of RF importance values")

        selected_set = set(np.asarray(self.selected_indices, dtype=int).tolist()) if self.selected_indices is not None else set()
        df = pd.DataFrame({
            'Original_Index': np.arange(len(feature_names), dtype=int),
            'Feature': list(feature_names),

            'RF_Importance': self.feature_importances,

            'RF_Gini_Importance': self.feature_importances
        })
        df = df.sort_values('RF_Gini_Importance', ascending=False).reset_index(drop=True)
        df['RF_Gini_Rank'] = np.arange(1, len(df) + 1, dtype=int)
        df['Cumulative_RF_Gini_Importance'] = df['RF_Gini_Importance'].cumsum()
        df['Selected'] = df['Original_Index'].isin(selected_set)
        return df

class RFESelector:

    def __init__(self, n_features_to_select=None, step=1, cv=3, random_state=42):
        self.n_features_to_select = n_features_to_select
        self.step = step
        self.cv = cv
        self.random_state = random_state
        self.rfe = None
        self.selected_indices = None
        self.selected_features = None
        self.support_ = None
        self.ranking_ = None
        self.initial_rf = None
        self.initial_rf_importances_ = None
        self.final_selected_rf_importances_ = None

    def fit(self, X, y, feature_names=None):

        self.initial_rf = RandomForestClassifier(
            n_estimators=100,
            random_state=self.random_state,
            n_jobs=-1
        )
        self.initial_rf.fit(X, y)
        self.initial_rf_importances_ = np.asarray(self.initial_rf.feature_importances_, dtype=float)

        base_estimator = RandomForestClassifier(
            n_estimators=100,
            random_state=self.random_state,
            n_jobs=-1
        )
        if self.n_features_to_select is None:
            temp_rf = RandomForestClassifier(n_estimators=50, random_state=self.random_state)
            temp_rf.fit(X, y)
            importances = temp_rf.feature_importances_
            indices = np.argsort(importances)[::-1]
            cumulative_importance = np.cumsum(importances[indices])
            self.n_features_to_select = np.argmax(cumulative_importance >= 0.95) + 1

        step = max(1, X.shape[1] // 20)
        self.rfe = RFECV(
            estimator=base_estimator,
            min_features_to_select=self.n_features_to_select,
            cv=self.cv,
            scoring='accuracy',
            n_jobs=-1,
            step=step
        )
        self.rfe.fit(X, y)
        self.support_ = np.asarray(self.rfe.support_, dtype=bool)
        self.ranking_ = np.asarray(self.rfe.ranking_, dtype=int)
        self.selected_indices = np.where(self.support_)[0]

        if feature_names is not None:
            self.selected_features = [feature_names[i] for i in self.selected_indices]
        else:
            self.selected_features = self.selected_indices

        final_estimator = getattr(self.rfe, 'estimator_', None)
        if final_estimator is not None and hasattr(final_estimator, 'feature_importances_'):
            final_imp = np.asarray(final_estimator.feature_importances_, dtype=float)
            if len(final_imp) == len(self.selected_indices):
                self.final_selected_rf_importances_ = final_imp
            else:
                self.final_selected_rf_importances_ = None
        return self

    def transform(self, X):
        if self.rfe is None:
            raise ValueError("Feature selector has not been fitted")
        return self.rfe.transform(X)

    def get_feature_importance_df(self, feature_names=None):

        if self.rfe is None:
            raise ValueError("The model has not been fitted")
        if feature_names is None:
            feature_names = [f'Feature_{i}' for i in range(len(self.ranking_))]
        if len(feature_names) != len(self.ranking_):
            raise ValueError("The number of feature_names does not match the number of RFE ranks")

        if self.initial_rf_importances_ is None:
            initial_imp = np.full(len(feature_names), np.nan)
        else:
            initial_imp = np.asarray(self.initial_rf_importances_, dtype=float)

        final_imp_all = np.full(len(feature_names), np.nan, dtype=float)
        if self.final_selected_rf_importances_ is not None:
            final_imp_all[self.selected_indices] = self.final_selected_rf_importances_

        df = pd.DataFrame({
            'Original_Index': np.arange(len(feature_names), dtype=int),
            'Feature': list(feature_names),
            'RFE_Ranking': self.ranking_,
            'Selected': self.support_,
            'Initial_RF_Gini_Importance': initial_imp,
            'Final_Selected_RF_Gini_Importance': final_imp_all
        })

        initial_order = np.argsort(-initial_imp)
        initial_rank = np.empty(len(initial_order), dtype=int)
        initial_rank[initial_order] = np.arange(1, len(initial_order) + 1, dtype=int)
        df['Initial_RF_Gini_Rank'] = initial_rank

        df['Final_Selected_RF_Gini_Rank'] = np.nan
        selected_df = df[df['Selected']].copy()
        selected_df = selected_df.sort_values('Final_Selected_RF_Gini_Importance', ascending=False)
        for rank, original_index in enumerate(selected_df['Original_Index'].tolist(), start=1):
            df.loc[df['Original_Index'] == original_index, 'Final_Selected_RF_Gini_Rank'] = rank

        df = df.sort_values(
            ['Selected', 'RFE_Ranking', 'Initial_RF_Gini_Importance'],
            ascending=[False, True, False]
        ).reset_index(drop=True)
        return df

    def get_cv_results_df(self):

        if self.rfe is None:
            raise ValueError("The model has not been fitted")
        cv_results = getattr(self.rfe, 'cv_results_', None)
        if cv_results is None:
            return pd.DataFrame()

        one_dimensional = {}
        for key, value in cv_results.items():
            arr = np.asarray(value)
            if arr.ndim == 1:
                one_dimensional[key] = arr
        if not one_dimensional:
            return pd.DataFrame()

        df = pd.DataFrame(one_dimensional)
        preferred = [c for c in ['n_features', 'mean_test_score', 'std_test_score', 'rank_test_score'] if c in df.columns]
        other = [c for c in df.columns if c not in preferred]
        return df[preferred + other]

def aggregate_feature_selection_importance(result_dir, method='rf', n_folds=5):
    fs_dir = os.path.join(result_dir, 'feature_selection_analysis')
    frames = []
    for fold in range(1, n_folds + 1):
        path = os.path.join(fs_dir, f'fold_{fold}_{method}_feature_importance.csv')
        if os.path.exists(path):
            df = pd.read_csv(path)
            df['Fold'] = fold
            frames.append(df)

    if not frames:
        return None

    all_df = pd.concat(frames, ignore_index=True)
    out_dir = os.path.join(result_dir, 'core_features')
    os.makedirs(out_dir, exist_ok=True)

    if method == 'rf':
        imp_col = 'RF_Gini_Importance' if 'RF_Gini_Importance' in all_df.columns else 'RF_Importance'
        agg = all_df.groupby('Feature', as_index=False).agg(
            Mean_RF_Gini_Importance=(imp_col, 'mean'),
            Std_RF_Gini_Importance=(imp_col, 'std'),
            Mean_RF_Gini_Rank=('RF_Gini_Rank', 'mean'),
            Selection_Count=('Selected', 'sum'),
            Fold_Count=('Fold', 'nunique')
        )
        agg['Selection_Frequency'] = agg['Selection_Count'] / agg['Fold_Count']
        agg = agg.sort_values(
            ['Selection_Frequency', 'Mean_RF_Gini_Importance'],
            ascending=[False, False]
        )
        out_csv = os.path.join(out_dir, 'rf_feature_importance_across_folds.csv')
        value_col = 'Mean_RF_Gini_Importance'
        title = 'RF Gini Importance across Folds'
    else:
        agg = all_df.groupby('Feature', as_index=False).agg(
            Mean_Initial_RF_Gini_Importance=('Initial_RF_Gini_Importance', 'mean'),
            Std_Initial_RF_Gini_Importance=('Initial_RF_Gini_Importance', 'std'),
            Mean_RFE_Ranking=('RFE_Ranking', 'mean'),
            Selection_Count=('Selected', 'sum'),
            Fold_Count=('Fold', 'nunique')
        )
        agg['Selection_Frequency'] = agg['Selection_Count'] / agg['Fold_Count']
        agg = agg.sort_values(
            ['Selection_Frequency', 'Mean_RFE_Ranking', 'Mean_Initial_RF_Gini_Importance'],
            ascending=[False, True, False]
        )
        out_csv = os.path.join(out_dir, 'rfe_feature_importance_across_folds.csv')
        value_col = 'Mean_Initial_RF_Gini_Importance'
        title = 'RFE Initial RF Gini Importance across Folds'

    agg.to_csv(out_csv, index=False, encoding='utf-8-sig')

    top = agg.sort_values(value_col, ascending=False).head(20)
    plt.figure(figsize=(12, 8))
    plt.barh(top['Feature'], top[value_col])
    plt.xlabel(value_col)
    plt.title(title)
    plt.gca().invert_yaxis()
    plt.tight_layout()
    out_png = out_csv.replace('.csv', '_top20.png')
    plt.savefig(out_png, dpi=300, bbox_inches='tight')
    plt.close()
    
    return agg

def get_rf_grid_search_params():
    return {
        'n_estimators': [100, 200, 300],
        'max_depth': [10, 20, 30, None],
        'min_samples_split': [2, 5, 10],
        'min_samples_leaf': [1, 2, 4],
        'max_features': ['sqrt', 'log2', 0.5]
    }

def get_mlp_grid_search_params():
    return {
        'hidden_layer_sizes': [(50,), (100,), (100, 50)],
        'activation': ['relu'],
        'alpha': [0.0001, 0.001, 0.01],
        'learning_rate_init': [0.001],
        'solver': ['adam'],
        'max_iter': [2000]
    }

def get_dt_grid_search_params():
    return {
        'criterion': ['gini', 'entropy'],
        'max_depth': [None, 10, 20, 30, 50],
        'min_samples_split': [2, 5, 10, 20],
        'min_samples_leaf': [1, 2, 4, 8],
        'max_features': [None, 'sqrt', 'log2']
    }

def get_svm_grid_search_params():
    return {
        'C': [0.1, 1, 10],
        'gamma': ['scale', 0.1, 1],
        'kernel': ['rbf'],
    }

def get_lgb_grid_search_params():
    return {
        'n_estimators': [100, 200, 300],
        'num_leaves': [31, 50, 70],
        'max_depth': [-1, 10, 20],
        'learning_rate': [0.01, 0.05, 0.1],
        'subsample': [0.8, 1.0],
        'colsample_bytree': [0.8, 1.0],
    }

def get_xgb_grid_search_params():
    return {
        'n_estimators': [100, 200, 300],
        'max_depth': [3, 6, 10],
        'learning_rate': [0.01, 0.05, 0.1],
        'subsample': [0.8, 1.0],
        'colsample_bytree': [0.8, 1.0],
        'reg_alpha': [0, 0.1, 1],
        'reg_lambda': [1, 1.5, 2],
    }

def train_classifier_with_grid_search(X_train, y_train, classifier_type, result_dir, fold_identifier=None):
    
    if classifier_type == 'rf':
        base = RandomForestClassifier(random_state=42, n_jobs=-1, class_weight='balanced')
        param_grid = get_rf_grid_search_params()
    elif classifier_type == 'mlp':
        base = MLPClassifier(max_iter=1000, random_state=42, early_stopping=True)
        param_grid = get_mlp_grid_search_params()
    elif classifier_type == 'dt':
        base = DecisionTreeClassifier(random_state=42, class_weight='balanced')
        param_grid = get_dt_grid_search_params()
    elif classifier_type == 'svm':
        base = SVC(probability=True, random_state=42, class_weight='balanced')
        param_grid = get_svm_grid_search_params()
        
    elif classifier_type == 'lgb':
        base = lgb.LGBMClassifier(random_state=42, n_jobs=-1, class_weight='balanced', verbose=-1)
        param_grid = get_lgb_grid_search_params()
    elif classifier_type == 'xgb':
        base = xgb.XGBClassifier(random_state=42, n_jobs=-1, use_label_encoder=False, eval_metric='logloss', verbosity=0)
        param_grid = get_xgb_grid_search_params()
    else:
        raise ValueError(f"Unknown classifier type: {classifier_type}")
    total_combinations = 1
    for v in param_grid.values():
        total_combinations *= len(v)
    
    grid_search = GridSearchCV(estimator=base, param_grid=param_grid, scoring='f1',
                               cv=3, n_jobs=-1, verbose=0, return_train_score=True)
    grid_search.fit(X_train, y_train)
    
    if result_dir:
        if fold_identifier is None:
            gs_dir = os.path.join(result_dir, "grid_search_details", "final_model")
            prefix = "final"
        else:
            gs_dir = os.path.join(result_dir, "grid_search_details", f"fold_{fold_identifier}")
            prefix = f"fold_{fold_identifier}"
        os.makedirs(gs_dir, exist_ok=True)
        pd.DataFrame(grid_search.cv_results_).to_csv(os.path.join(gs_dir, f"{prefix}_grid_search_full_results.csv"), index=False)
        with open(os.path.join(gs_dir, f"{prefix}_best_parameters.json"), 'w') as f:
            json.dump({'best_params': grid_search.best_params_, 'best_score': float(grid_search.best_score_)}, f, indent=4)
    return grid_search.best_estimator_, grid_search.best_params_

def plot_learning_curve(model, X_train, y_train, title, output_path, cv=5, scoring='accuracy'):
    try:
        train_sizes, train_scores, test_scores = learning_curve(
            model, X_train, y_train, cv=cv, scoring=scoring,
            train_sizes=np.linspace(0.1, 1.0, 10), n_jobs=-1)
        train_mean = np.mean(train_scores, axis=1)
        train_std = np.std(train_scores, axis=1)
        test_mean = np.mean(test_scores, axis=1)
        test_std = np.std(test_scores, axis=1)
        plt.figure(figsize=(10, 6))
        plt.fill_between(train_sizes, train_mean - train_std, train_mean + train_std, alpha=0.1, color="blue")
        plt.fill_between(train_sizes, test_mean - test_std, test_mean + test_std, alpha=0.1, color="orange")
        plt.plot(train_sizes, train_mean, 'o-', color="blue", label="Training score")
        plt.plot(train_sizes, test_mean, 'o-', color="orange", label="Cross-validation score")
        plt.xlabel("Training examples")
        plt.ylabel(scoring)
        plt.title(title)
        plt.legend(loc="best")
        plt.grid(True)
        plt.tight_layout()
        plt.savefig(output_path, dpi=300)
        plt.close()
        
        return True
    except Exception as e:
        print(f"Failed to plot learning curve: {str(e)}")
        return False

def plot_mlp_loss_curve(mlp_model, output_path):
    if hasattr(mlp_model, 'loss_curve_') and mlp_model.loss_curve_:
        plt.figure(figsize=(10, 6))
        plt.plot(mlp_model.loss_curve_, 'b-')
        plt.xlabel('Iteration')
        plt.ylabel('Loss')
        plt.title('MLP Training Loss Curve')
        plt.grid(True)
        plt.tight_layout()
        plt.savefig(output_path, dpi=300)
        plt.close()
        
        return True
    return False

def plot_validation_curve(mlp_model, output_path):
    if hasattr(mlp_model, 'validation_scores_') and mlp_model.validation_scores_:
        plt.figure(figsize=(10, 6))
        plt.plot(mlp_model.validation_scores_, 'r-')
        plt.xlabel('Iteration')
        plt.ylabel('Validation Score')
        plt.title('MLP Validation Score Curve')
        plt.grid(True)
        plt.tight_layout()
        plt.savefig(output_path, dpi=300)
        plt.close()
        
        return True
    return False

def plot_feature_importance(importance_df, filename, metric_dir, importance_col='Importance'):
    if importance_col not in importance_df.columns:
        candidates = ['Importance', 'RF_Importance', 'RFE_Ranking', 'Importance_Score']
        for col in candidates:
            if col in importance_df.columns:
                importance_col = col
                break
        else:
            num_cols = importance_df.select_dtypes(include=[np.number]).columns
            if len(num_cols) > 0:
                importance_col = num_cols[0]
            else:
                print(f"Warning: no numeric column found; skipping plot {filename}")
                return
    if 'RFE_Ranking' in importance_col:
        sorted_df = importance_df.nsmallest(20, importance_col)
        plt.figure(figsize=(12, 8))
        plt.barh(sorted_df['Feature'], sorted_df[importance_col], color='steelblue')
        plt.xlabel('RFE Ranking (lower is better)')
    else:
        sorted_df = importance_df.nlargest(20, importance_col)
        plt.figure(figsize=(12, 8))
        plt.barh(sorted_df['Feature'], sorted_df[importance_col], color='steelblue')
        plt.xlabel('Importance')
    plt.title(f'Feature Importance - {filename}')
    plt.gca().invert_yaxis()
    plt.tight_layout()
    out_path = os.path.join(metric_dir, "feature_importance", f"{filename}.png")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    plt.savefig(out_path, dpi=300)
    plt.close()
    

def plot_smote_distribution(y_before, y_after, filename, metric_dir):
    before_counts = np.bincount(y_before)
    after_counts = np.bincount(y_after)
    classes = np.arange(max(len(before_counts), len(after_counts)))
    before_counts = np.pad(before_counts, (0, len(classes) - len(before_counts)), 'constant')
    after_counts = np.pad(after_counts, (0, len(classes) - len(after_counts)), 'constant')
    plt.figure(figsize=(10, 6))
    bar_width = 0.35
    x = np.arange(len(classes))
    plt.bar(x - bar_width / 2, before_counts, bar_width, label='Before SMOTE', color='skyblue')
    plt.bar(x + bar_width / 2, after_counts, bar_width, label='After SMOTE', color='coral')
    plt.xlabel('Class')
    plt.ylabel('Count')
    plt.title(f'SMOTE Class Distribution - {filename}')
    plt.xticks(x, classes)
    plt.legend()
    plt.grid(axis='y', alpha=0.3)
    plt.tight_layout()
    out_path = os.path.join(metric_dir, "smote_distributions", f"{filename}.png")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    plt.savefig(out_path, dpi=300)
    plt.close()
    

def kfold_cross_validation(X, y, feature_names, config, result_dir, metric_dir, threshold=0.5):
    n_folds = config['n_folds']
    feature_selector_type = config['feature_selector']
    classifier_type = config['classifier']
    perform_grid_search = config.get('perform_grid_search', True)
    model_key = f"{feature_selector_type}+{classifier_type}"

    
    
    
    

    for d in ["models", "fold_visualizations", "results", "roc_curves",
              "feature_importance", "confusion_matrices", "core_features",
              "pca_analysis", "feature_selection_analysis", "fold_details_data",
              "grid_search_details", "model_comparison"]:
        os.makedirs(os.path.join(result_dir, d), exist_ok=True)

    kfold = StratifiedKFold(n_splits=n_folds, shuffle=True, random_state=42)
    fold_accuracies, fold_auc_scores, all_results = [], [], []

    for fold, (train_idx, val_idx) in enumerate(kfold.split(X, y)):
        
        X_train, X_val = X[train_idx], X[val_idx]
        y_train, y_val = y[train_idx], y[val_idx]

        fold_dir = os.path.join(result_dir, "fold_details_data", f"fold_{fold + 1}")
        os.makedirs(fold_dir, exist_ok=True)
        pd.DataFrame(X_train, columns=feature_names).assign(Label=y_train).to_csv(
            os.path.join(fold_dir, f'fold_{fold + 1}_train_set.csv'), index=False)
        pd.DataFrame(X_val, columns=feature_names).assign(Label=y_val).to_csv(
            os.path.join(fold_dir, f'fold_{fold + 1}_validation_set.csv'), index=False)
        

        if feature_selector_type == 'rf':
            selector = RFSelector(n_estimators=100, random_state=42)
            selector.fit(X_train, y_train, feature_names)
            X_train_selected = selector.transform(X_train)
            X_val_selected = selector.transform(X_val)
            selected_features = selector.selected_features

            imp_df = selector.get_feature_importance_df(feature_names)
            rf_imp_path = os.path.join(
                result_dir, "feature_selection_analysis",
                f"fold_{fold + 1}_rf_feature_importance.csv"
            )
            imp_df.to_csv(rf_imp_path, index=False, encoding='utf-8-sig')
            plot_feature_importance(
                imp_df, f"{model_key}_fold{fold + 1}_rf_gini_importance",
                metric_dir, importance_col='RF_Gini_Importance'
            )
        else:
            selector = RFESelector(n_features_to_select=None, step=1, cv=3, random_state=42)
            selector.fit(X_train, y_train, feature_names)
            X_train_selected = selector.transform(X_train)
            X_val_selected = selector.transform(X_val)
            selected_features = selector.selected_features

            imp_df = selector.get_feature_importance_df(feature_names)
            rfe_imp_path = os.path.join(
                result_dir, "feature_selection_analysis",
                f"fold_{fold + 1}_rfe_feature_importance.csv"
            )
            imp_df.to_csv(rfe_imp_path, index=False, encoding='utf-8-sig')
            plot_feature_importance(
                imp_df, f"{model_key}_fold{fold + 1}_rfe_ranking",
                metric_dir, importance_col='RFE_Ranking'
            )
            plot_feature_importance(
                imp_df, f"{model_key}_fold{fold + 1}_rfe_initial_rf_gini",
                metric_dir, importance_col='Initial_RF_Gini_Importance'
            )
            final_selected_imp = imp_df[
                imp_df['Selected'] & imp_df['Final_Selected_RF_Gini_Importance'].notna()
            ].copy()
            if not final_selected_imp.empty:
                plot_feature_importance(
                    final_selected_imp,
                    f"{model_key}_fold{fold + 1}_rfe_final_selected_rf_gini",
                    metric_dir,
                    importance_col='Final_Selected_RF_Gini_Importance'
                )

            rfe_cv_df = selector.get_cv_results_df()
            if not rfe_cv_df.empty:
                rfe_cv_df.to_csv(
                    os.path.join(
                        result_dir, "feature_selection_analysis",
                        f"fold_{fold + 1}_rfe_rfecv_cv_results.csv"
                    ),
                    index=False, encoding='utf-8-sig'
                )

        pd.DataFrame(X_train_selected, columns=selected_features).assign(Label=y_train).to_csv(
            os.path.join(fold_dir, f'fold_{fold + 1}_train_set_selected.csv'), index=False)
        

        pca_analyzer = PCAAnalyzer(variance_threshold=0.95, random_state=42)
        X_train_pca = pca_analyzer.fit_transform(X_train_selected, selected_features, y_train)
        X_val_pca = pca_analyzer.transform(X_val_selected)
        pca_columns = [f'PC{i + 1}' for i in range(pca_analyzer.n_components)]
        pca_data_path = os.path.join(metric_dir, "pca_data", f"{model_key}_fold{fold + 1}_train_pca.csv")
        os.makedirs(os.path.dirname(pca_data_path), exist_ok=True)
        pd.DataFrame(X_train_pca).to_csv(pca_data_path, index=False)
        pca_analyzer.save_analysis(result_dir, prefix=f"fold_{fold + 1}_")

        try:
            min_class = min(np.unique(y_train, return_counts=True)[1])
            k_neighbors = min(5, min_class - 1) if min_class > 1 else 1
            smote = BorderlineSMOTE(random_state=42, kind='borderline-1')
            X_train_res, y_train_res = smote.fit_resample(X_train_pca, y_train)
            plot_smote_distribution(y_train, y_train_res, f"{model_key}_fold{fold + 1}_smote", metric_dir)
            excel_path = os.path.join(metric_dir, "smote_distributions", f"{model_key}_fold{fold + 1}_smote.xlsx")
            save_smote_distribution_excel(y_train, y_train_res, excel_path, f"SMOTE Change - Fold {fold + 1}")
        except Exception as e:
            print(f"Fold {fold + 1} SMOTE failed: {str(e)}. Using the original data.")
            X_train_res, y_train_res = X_train_pca, y_train

        baseline_params = {
            'rf': {'n_estimators': 100, 'max_depth': None, 'min_samples_split': 2, 'min_samples_leaf': 1,
                   'max_features': 'sqrt'},
            'mlp': {'hidden_layer_sizes': (100,), 'activation': 'relu', 'solver': 'adam', 'alpha': 0.0001,
                    'learning_rate_init': 0.001},
            'dt': {'criterion': 'gini', 'max_depth': None, 'min_samples_split': 2, 'min_samples_leaf': 1,
                   'max_features': None},
            'svm': {'C': 1.0, 'kernel': 'rbf', 'gamma': 'scale'},
            'lgb': {'n_estimators': 100, 'num_leaves': 31, 'max_depth': -1, 'learning_rate': 0.1,
                    'subsample': 1.0, 'colsample_bytree': 1.0},
            'xgb': {'n_estimators': 100, 'max_depth': 6, 'learning_rate': 0.1,
                    'subsample': 1.0, 'colsample_bytree': 1.0, 'reg_alpha': 0, 'reg_lambda': 1}
        }[classifier_type]

        if classifier_type == 'rf':
            baseline_model = RandomForestClassifier(**baseline_params, random_state=42, n_jobs=-1, class_weight='balanced')
        elif classifier_type == 'mlp':
            baseline_model = MLPClassifier(**baseline_params, max_iter=1000, random_state=42, early_stopping=True)
        elif classifier_type == 'dt':
            baseline_model = DecisionTreeClassifier(**baseline_params, random_state=42, class_weight='balanced')
        elif classifier_type == 'svm':
            baseline_model = SVC(**baseline_params, probability=True, random_state=42, class_weight='balanced')
        elif classifier_type == 'lgb':
            baseline_model = lgb.LGBMClassifier(**baseline_params, random_state=42, n_jobs=-1, class_weight='balanced', verbose=-1)
        elif classifier_type == 'xgb':
            baseline_model = xgb.XGBClassifier(**baseline_params, random_state=42, n_jobs=-1,
                                               use_label_encoder=False, eval_metric='logloss', verbosity=0)
        else:
            raise ValueError(f"Unknown classifier: {classifier_type}")

        baseline_model.fit(X_train_res, y_train_res)
        baseline_acc = accuracy_score(y_val, baseline_model.predict(X_val_pca))
        baseline_result = {'fold': fold + 1, 'model_type': 'baseline', 'accuracy': float(baseline_acc),
                           'parameters': baseline_params}

        if perform_grid_search:
            optimized_model, best_params = train_classifier_with_grid_search(X_train_res, y_train_res, classifier_type,
                                                                             result_dir, fold + 1)
            optimized_acc = accuracy_score(y_val, optimized_model.predict(X_val_pca))
            optimized_result = {'fold': fold + 1, 'model_type': 'optimized', 'accuracy': float(optimized_acc),
                                'parameters': best_params, 'best_cv_score': None}
            model = optimized_model
        else:
            model = baseline_model
            optimized_result = None

        val_probs_full = model.predict_proba(X_val_pca)
        val_probs = val_probs_full[:, 1]
        val_preds = (val_probs >= threshold).astype(int)
        fold_acc = accuracy_score(y_val, val_preds)
        fold_accuracies.append(fold_acc)
        if len(np.unique(y_val)) > 1:
            fold_auc = calculate_binary_auc(y_val, val_probs_full, model.classes_, fold + 1)
            fold_auc_scores.append(fold_auc)
            roc_path_model = os.path.join(result_dir, "roc_curves", f"fold_{fold + 1}_validation_roc.png")
            if plot_roc_curve_binary(y_val, val_probs_full, model.classes_,
                                     f'ROC Curve - Validation Set (Fold {fold + 1})', roc_path_model):
                roc_path_metric = os.path.join(metric_dir, "roc_curves", f"{model_key}_fold{fold + 1}_validation_roc.png")
                if os.path.exists(roc_path_model):
                    shutil.copy2(roc_path_model, roc_path_metric)
            else:
                print(f" Fold {fold+1} ROC plot failed; skipping copy")
        else:
            fold_auc_scores.append(0.0)

        unique_cls = np.unique(np.concatenate([y_val, val_preds]))
        class_labels = [str(c) for c in unique_cls]
        save_confusion_matrix(y_val, val_preds,
                              os.path.join(result_dir, "confusion_matrices", f"fold_{fold + 1}_validation_cm.png"),
                              classes=class_labels, title=f"Validation CM - Fold {fold + 1}")
        save_confusion_matrix(y_val, val_preds,
                              os.path.join(result_dir, "confusion_matrices", f"fold_{fold + 1}_validation_cm_norm.png"),
                              classes=class_labels, normalize=True, title=f"Validation CM (Norm) - Fold {fold + 1}")
        cm_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_fold{fold + 1}_validation_cm.png")
        save_confusion_matrix(y_val, val_preds, cm_path, classes=class_labels, title=f"Validation CM - Fold {fold + 1}")
        cm_norm_path = os.path.join(metric_dir, "confusion_matrices",
                                    f"{model_key}_fold{fold + 1}_validation_cm_norm.png")
        save_confusion_matrix(y_val, val_preds, cm_norm_path, classes=class_labels, normalize=True,
                              title=f"Validation CM (Norm) - Fold {fold + 1}")

        fold_model_dir = os.path.join(result_dir, "models", f'fold_{fold + 1}')
        os.makedirs(fold_model_dir, exist_ok=True)
        joblib.dump(model, os.path.join(fold_model_dir, f'{classifier_type}_model.joblib'))
        joblib.dump(pca_analyzer, os.path.join(fold_model_dir, 'pca_analyzer.joblib'))
        joblib.dump(selector, os.path.join(fold_model_dir, 'feature_selector.joblib'))

        fold_result = {
            'fold': fold + 1,
            'baseline_accuracy': float(baseline_acc),
            'optimized_accuracy': float(optimized_acc) if optimized_result else None,
            'accuracy': float(optimized_acc) if optimized_result else float(baseline_acc),
            'auc': float(fold_auc_scores[-1]),
            'selected_features': selected_features,
            'num_selected_features': len(selected_features),
            'pca_components': pca_analyzer.n_components,
            'pca_variance_explained': float(sum(pca_analyzer.pca.explained_variance_ratio_)),
            'use_grid_search': perform_grid_search
        }
        if classifier_type == 'rf':
            fold_result['rf_selected_features'] = selected_features
        else:
            fold_result['rfe_selected_features'] = selected_features
        all_results.append(fold_result)

        

    with open(os.path.join(result_dir, 'results/kfold_results.json'), 'w') as f:
        json.dump(convert_types(all_results), f, indent=4)

    mean_acc = np.mean(fold_accuracies) if fold_accuracies else 0.0
    std_acc = np.std(fold_accuracies) if fold_accuracies else 0.0
    mean_auc = np.mean(fold_auc_scores) if fold_auc_scores else 0.0
    std_auc = np.std(fold_auc_scores) if fold_auc_scores else 0.0
    print(f"\nMean accuracy: {mean_acc:.4f} ± {std_acc:.4f}")
    print(f"Mean AUC: {mean_auc:.4f} ± {std_auc:.4f}")

    pd.DataFrame({'Fold': range(1, n_folds + 1), 'Accuracy': fold_accuracies, 'AUC': fold_auc_scores}).to_csv(
        os.path.join(result_dir, 'results/kfold_cv_results.csv'), index=False)

    core_features = analyze_feature_selection_consistency(all_results, result_dir, method=feature_selector_type)

    aggregate_feature_selection_importance(
        result_dir,
        method=feature_selector_type,
        n_folds=n_folds
    )

    baseline_accs = [r['baseline_accuracy'] for r in all_results]
    optimized_accs = [r['optimized_accuracy'] for r in all_results if r['optimized_accuracy'] is not None]
    if optimized_accs:
        
        print(f"   Mean baseline accuracy: {np.mean(baseline_accs):.4f}")
        print(f"   Mean optimized accuracy: {np.mean(optimized_accs):.4f}")
        print(f"   Accuracy improvement: {np.mean(optimized_accs) - np.mean(baseline_accs):.4f}")
    return fold_accuracies, mean_acc, core_features, all_results, fold_auc_scores

def train_final_model(X_train, y_train, core_features, feature_names, config, result_dir, metric_dir):
    feature_selector_type = config['feature_selector']
    classifier_type = config['classifier']
    perform_grid_search = config.get('perform_grid_search', True)
    model_key = f"{feature_selector_type}+{classifier_type}"

    
    
    

    dirs = ["models", "data_analysis", "roc_curves", "feature_importance",
            "confusion_matrices", "pca_analysis", "feature_selection_analysis", "final_model_details"]
    for d in dirs:
        os.makedirs(os.path.join(result_dir, d), exist_ok=True)
    final_dir = os.path.join(result_dir, "final_model_details")
    os.makedirs(final_dir, exist_ok=True)

    core_indices = [feature_names.index(feat) for feat in core_features if feat in feature_names]
    X_train_core = X_train[:, core_indices]
    

    if feature_selector_type == 'rf':
        selector = RFSelector(n_estimators=100, random_state=42)
        selector.fit(X_train_core, y_train, core_features)
        X_train_selected = selector.transform(X_train_core)
        selected_features = selector.selected_features

        final_imp_df = selector.get_feature_importance_df(core_features)
        final_imp_df.to_csv(
            os.path.join(result_dir, "feature_selection_analysis", "final_rf_feature_importance.csv"),
            index=False, encoding='utf-8-sig'
        )
        plot_feature_importance(
            final_imp_df, f"{model_key}_final_rf_gini_importance",
            metric_dir, importance_col='RF_Gini_Importance'
        )
    else:
        selector = RFESelector(n_features_to_select=None, step=1, cv=3, random_state=42)
        selector.fit(X_train_core, y_train, core_features)
        X_train_selected = selector.transform(X_train_core)
        selected_features = selector.selected_features

        final_imp_df = selector.get_feature_importance_df(core_features)
        final_imp_df.to_csv(
            os.path.join(result_dir, "feature_selection_analysis", "final_rfe_feature_importance.csv"),
            index=False, encoding='utf-8-sig'
        )
        plot_feature_importance(
            final_imp_df, f"{model_key}_final_rfe_ranking",
            metric_dir, importance_col='RFE_Ranking'
        )
        plot_feature_importance(
            final_imp_df, f"{model_key}_final_rfe_initial_rf_gini",
            metric_dir, importance_col='Initial_RF_Gini_Importance'
        )
        final_selected_imp = final_imp_df[
            final_imp_df['Selected'] & final_imp_df['Final_Selected_RF_Gini_Importance'].notna()
        ].copy()
        if not final_selected_imp.empty:
            plot_feature_importance(
                final_selected_imp, f"{model_key}_final_rfe_selected_rf_gini",
                metric_dir, importance_col='Final_Selected_RF_Gini_Importance'
            )

        final_rfe_cv_df = selector.get_cv_results_df()
        if not final_rfe_cv_df.empty:
            final_rfe_cv_df.to_csv(
                os.path.join(
                    result_dir, "feature_selection_analysis",
                    "final_rfe_rfecv_cv_results.csv"
                ),
                index=False, encoding='utf-8-sig'
            )

    pd.DataFrame(X_train_selected, columns=selected_features).assign(Label=y_train).to_csv(
        os.path.join(final_dir, 'final_train_set_selected.csv'), index=False)
    print(f"The final model retained {len(selected_features)} features after feature selection.")

    pca_analyzer = PCAAnalyzer(variance_threshold=0.95, random_state=42)
    X_train_pca = pca_analyzer.fit_transform(X_train_selected, selected_features, y_train)
    pca_columns = [f'PC{i + 1}' for i in range(pca_analyzer.n_components)]
    pd.DataFrame(X_train_pca, columns=pca_columns).assign(Label=y_train).to_csv(
        os.path.join(final_dir, 'final_train_set_pca.csv'), index=False)
    print(f"PCA retained {pca_analyzer.n_components} principal components.")
    pca_analyzer.save_analysis(result_dir, prefix="final_")

    try:
        min_class = min(np.unique(y_train, return_counts=True)[1])
        k = min(5, min_class - 1) if min_class > 1 else 1
        smote = BorderlineSMOTE(random_state=42, kind='borderline-1')
        X_train_res, y_train_res = smote.fit_resample(X_train_pca, y_train)
        pd.DataFrame(X_train_res, columns=pca_columns).assign(Label=y_train_res).to_csv(
            os.path.join(final_dir, 'final_train_set_smote.csv'), index=False)
        plot_smote_distribution(y_train, y_train_res, f"{model_key}_final_smote", metric_dir)
        excel_path = os.path.join(metric_dir, "smote_distributions", f"{model_key}_final_smote.xlsx")
        save_smote_distribution_excel(y_train, y_train_res, excel_path, "Final Model SMOTE Change")
    except Exception as e:
        print(f"SMOTE failed: {str(e)}. Using the original data.")
        X_train_res, y_train_res = X_train_pca, y_train

    if perform_grid_search:
        
        model, best_params = train_classifier_with_grid_search(X_train_res, y_train_res, classifier_type, result_dir,
                                                               "final")
        if classifier_type == 'mlp' and hasattr(model, 'loss_curve_'):
            loss_path = os.path.join(metric_dir, "learning_curves", f"{model_key}_final_mlp_loss.png")
            plot_mlp_loss_curve(model, loss_path)
            val_path = os.path.join(metric_dir, "learning_curves", f"{model_key}_final_mlp_validation.png")
            plot_validation_curve(model, val_path)
    else:

        if classifier_type == 'rf':
            model = RandomForestClassifier(n_estimators=100, random_state=42, n_jobs=-1, class_weight='balanced')
        elif classifier_type == 'mlp':
            model = MLPClassifier(hidden_layer_sizes=(100,), activation='relu', solver='adam',
                                  alpha=0.0001, learning_rate_init=0.001, max_iter=1000, random_state=42,
                                  early_stopping=True)
        elif classifier_type == 'dt':
            model = DecisionTreeClassifier(random_state=42, class_weight='balanced')
        elif classifier_type == 'svm':
            model = SVC(C=1.0, kernel='rbf', gamma='scale', probability=True, random_state=42, class_weight='balanced')
        elif classifier_type == 'lgb':
            model = lgb.LGBMClassifier(n_estimators=100, num_leaves=31, max_depth=-1, learning_rate=0.1,
                                       subsample=1.0, colsample_bytree=1.0, random_state=42, n_jobs=-1,
                                       class_weight='balanced', verbose=-1)
        elif classifier_type == 'xgb':
            model = xgb.XGBClassifier(n_estimators=100, max_depth=6, learning_rate=0.1,
                                      subsample=1.0, colsample_bytree=1.0, reg_alpha=0, reg_lambda=1,
                                      random_state=42, n_jobs=-1, use_label_encoder=False, eval_metric='logloss',
                                      verbosity=0)
        else:
            raise ValueError(f"Unknown classifier: {classifier_type}")
        model.fit(X_train_res, y_train_res)
        best_params = {}
        if classifier_type == 'mlp' and hasattr(model, 'loss_curve_'):
            loss_path = os.path.join(metric_dir, "learning_curves", f"{model_key}_final_mlp_loss.png")
            plot_mlp_loss_curve(model, loss_path)

    train_preds_smote = model.predict(X_train_res)
    train_probs_smote_full = model.predict_proba(X_train_res)

    train_preds_original = model.predict(X_train_pca)
    train_probs_original_full = model.predict_proba(X_train_pca)

    if 1 in list(model.classes_):
        positive_index = list(model.classes_).index(1)
    else:
        positive_index = 1 if len(model.classes_) > 1 else 0

    train_probs_original = train_probs_original_full[:, positive_index]
    train_metrics = compute_binary_metrics(
        y_train,
        train_preds_original,
        train_probs_original
    )
    train_accuracy = train_metrics['Accuracy']

    print(
        f"Final-model metrics on the original training set: "
        f"Accuracy={train_metrics['Accuracy']:.4f}, "
        f"BA={train_metrics['BA']:.4f}, "
        f"G-mean={train_metrics['G-mean']:.4f}, "
        f"F1={train_metrics['F1-score']:.4f}, "
        f"Recall={train_metrics['Recall']:.4f}, "
        f"ROC-AUC={train_metrics['ROC-AUC']:.4f}"
    )

    train_metrics_path = os.path.join(final_dir, 'final_train_metrics.json')
    with open(train_metrics_path, 'w', encoding='utf-8') as f:
        json.dump(convert_types({
            **train_metrics,
            'training_size_original': int(len(y_train)),
            'training_size_smote': int(len(y_train_res)),
            'evaluation_dataset': 'original_training_set_non_resampled',
            'model_fit_dataset': 'BorderlineSMOTE_resampled_training_set'
        }), f, indent=4, ensure_ascii=False)

    lc_path = os.path.join(metric_dir, "learning_curves", f"{model_key}_final_learning_curve.png")
    plot_learning_curve(model, X_train_res, y_train_res, f'Learning Curve - {model_key}', lc_path)

    joblib.dump(model, os.path.join(result_dir, "models", f"final_{classifier_type}_model.joblib"))
    joblib.dump(pca_analyzer, os.path.join(result_dir, "models", "final_pca_analyzer.joblib"))
    joblib.dump(selector, os.path.join(result_dir, "models", "final_feature_selector.joblib"))
    with open(os.path.join(result_dir, "models", "final_model_info.json"), 'w', encoding='utf-8') as f:
        json.dump(convert_types({
            'core_features': core_features,
            'selected_features': selected_features,
            'num_original_features': len(feature_names),
            'num_core_features': len(core_features),
            'num_selected_features': len(selected_features),
            'best_parameters': best_params,
            'training_accuracy': float(train_accuracy),
            'training_metrics_original': train_metrics,
            'training_size_original': int(len(y_train)),
            'training_size_smote': int(len(y_train_res))
        }), f, indent=4, ensure_ascii=False)

    try:
        roc_model_path = os.path.join(result_dir, "roc_curves", "final_model_training_roc.png")
        plot_roc_curve_binary(
            y_train,
            train_probs_original_full,
            model.classes_,
            'ROC Curve - Final Model (Original Training Set)',
            roc_model_path
        )
        roc_metric_path = os.path.join(metric_dir, "roc_curves", f"{model_key}_final_training_roc.png")
        shutil.copy2(roc_model_path, roc_metric_path)
        roc_data_src = roc_model_path.replace('.png', '_data.csv')
        roc_data_dst = roc_metric_path.replace('.png', '_data.csv')
        if os.path.exists(roc_data_src):
            shutil.copy2(roc_data_src, roc_data_dst)
    except Exception as e:
        print(f"Failed to plot the final model ROC curve for the original training set: {str(e)}")

    unique_cls_original = np.unique(y_train)
    class_labels_original = [str(c) for c in unique_cls_original]

    save_confusion_matrix(
        y_train, train_preds_original,
        os.path.join(result_dir, "confusion_matrices", "final_model_training_cm.png"),
        classes=class_labels_original,
        title="Final Model - Original Training CM"
    )
    save_confusion_matrix(
        y_train, train_preds_original,
        os.path.join(result_dir, "confusion_matrices", "final_model_training_cm_norm.png"),
        classes=class_labels_original,
        normalize=True,
        title="Final Model - Original Training CM (Norm)"
    )

    cm_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_final_training_cm.png")
    save_confusion_matrix(
        y_train, train_preds_original, cm_path,
        classes=class_labels_original,
        title="Final Model - Original Training CM"
    )
    cm_norm_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_final_training_cm_norm.png")
    save_confusion_matrix(
        y_train, train_preds_original, cm_norm_path,
        classes=class_labels_original,
        normalize=True,
        title="Final Model - Original Training CM (Norm)"
    )

    excel_path = os.path.join(result_dir, "confusion_matrices", "final_model_training_cm.xlsx")
    save_confusion_matrix_excel(
        y_train, train_preds_original, excel_path,
        classes=class_labels_original,
        normalize=False
    )
    excel_norm_path = os.path.join(result_dir, "confusion_matrices", "final_model_training_cm_norm.xlsx")
    save_confusion_matrix_excel(
        y_train, train_preds_original, excel_norm_path,
        classes=class_labels_original,
        normalize=True
    )

    metric_excel_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_final_training_cm.xlsx")
    save_confusion_matrix_excel(
        y_train, train_preds_original, metric_excel_path,
        classes=class_labels_original,
        normalize=False
    )
    metric_excel_norm_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_final_training_cm_norm.xlsx")
    save_confusion_matrix_excel(
        y_train, train_preds_original, metric_excel_norm_path,
        classes=class_labels_original,
        normalize=True
    )

    try:
        smote_roc_path = os.path.join(result_dir, "roc_curves", "final_model_training_smote_roc.png")
        plot_roc_curve_binary(
            y_train_res,
            train_probs_smote_full,
            model.classes_,
            'ROC Curve - Final Model (Training SMOTE, Diagnostic Only)',
            smote_roc_path
        )
    except Exception as e:
        print(f"Failed to plot the diagnostic ROC curve for the SMOTE-resampled training set: {str(e)}")

    unique_cls_smote = np.unique(y_train_res)
    class_labels_smote = [str(c) for c in unique_cls_smote]
    save_confusion_matrix(
        y_train_res, train_preds_smote,
        os.path.join(result_dir, "confusion_matrices", "final_model_training_smote_cm.png"),
        classes=class_labels_smote,
        title="Final Model - Training SMOTE CM (Diagnostic Only)"
    )
    save_confusion_matrix_excel(
        y_train_res, train_preds_smote,
        os.path.join(result_dir, "confusion_matrices", "final_model_training_smote_cm.xlsx"),
        classes=class_labels_smote,
        normalize=False
    )

    raw_perm_imp = None
    return model, pca_analyzer, selector, core_indices, train_accuracy, raw_perm_imp, selected_features

def analyze_factor_importance(model, pca_analyzer, X_train_original, y_train_original,
                              X_train_pca_smoted, y_train_smoted, selected_features, all_feature_names,
                              classifier_type, result_dir, metric_dir, model_key, n_permutations=30):
    
    
    fa_dir = os.path.join(result_dir, "factor_analysis")
    os.makedirs(fa_dir, exist_ok=True)

    pca_cols = [f'PC{i + 1}' for i in range(X_train_pca_smoted.shape[1])]

    raw_perm_importance = None
    
    try:
        if hasattr(model, 'get_params'):
            model_params = model.get_params()
        else:
            model_params = {}
        params = model_params.copy()
        conflict_keys = ['random_state', 'n_jobs', 'max_iter', 'early_stopping', 'probability', 'verbose', 'verbosity',
                         'use_label_encoder', 'eval_metric']
        for k in conflict_keys:
            params.pop(k, None)

        if classifier_type == 'rf':
            raw_model = RandomForestClassifier(**params, random_state=42, n_jobs=-1)
        elif classifier_type == 'mlp':
            raw_model = MLPClassifier(**params, random_state=42, max_iter=1000, early_stopping=True)
        elif classifier_type == 'dt':
            raw_model = DecisionTreeClassifier(**params, random_state=42)
        elif classifier_type == 'svm':
            raw_model = SVC(**params, probability=True, random_state=42)
        elif classifier_type == 'lgb':
            raw_model = lgb.LGBMClassifier(**params, random_state=42, n_jobs=-1, verbose=-1)
        elif classifier_type == 'xgb':
            raw_model = xgb.XGBClassifier(**params, random_state=42, n_jobs=-1,
                                          use_label_encoder=False, eval_metric='logloss', verbosity=0)
        else:
            raw_model = None

        if raw_model is not None:
            raw_model.fit(X_train_original, y_train_original)
            perm_result = permutation_importance(
                raw_model, X_train_original, y_train_original,
                n_repeats=n_permutations, random_state=42, n_jobs=-1, scoring='roc_auc'
            )
            raw_perm_importance = pd.DataFrame({
                'Feature': selected_features,
                'Permutation_Importance_Mean': perm_result.importances_mean,
                'Permutation_Importance_Std': perm_result.importances_std
            }).sort_values('Permutation_Importance_Mean', ascending=False)

            raw_perm_path = os.path.join(fa_dir, f'{classifier_type}_raw_permutation_importance.csv')
            raw_perm_importance.to_csv(raw_perm_path, index=False)
            metric_perm_dir = os.path.join(metric_dir, "permutation_importance")
            os.makedirs(metric_perm_dir, exist_ok=True)
            metric_perm_path = os.path.join(metric_perm_dir, f"{model_key}_raw_permutation_importance.csv")
            raw_perm_importance.to_csv(metric_perm_path, index=False)
            

            plt.figure(figsize=(12, 8))
            top20 = raw_perm_importance.head(20)
            plt.barh(top20['Feature'], top20['Permutation_Importance_Mean'],
                     xerr=top20['Permutation_Importance_Std'], capsize=3)
            plt.xlabel('Permutation Importance (AUC drop)')
            plt.title(f'Raw Feature Permutation Importance - {classifier_type.upper()}')
            plt.gca().invert_yaxis()
            plt.tight_layout()
            img_path = os.path.join(metric_perm_dir, f"{model_key}_raw_permutation_importance.png")
            plt.savefig(img_path, dpi=300, bbox_inches='tight')
            plt.close()
        else:
            print("Could not create the raw-feature model; skipping permutation-importance calculation.")
    except Exception as e:
        print(f"Failed to calculate permutation importance for original features: {str(e)}")
        traceback.print_exc()

    raw_self_importance = None
    if raw_model is not None:
        try:
            if classifier_type in ['rf', 'dt', 'lgb', 'xgb'] and hasattr(raw_model, 'feature_importances_'):
                self_imp = raw_model.feature_importances_
                if len(self_imp) == len(selected_features):
                    imp_name = f'{classifier_type.upper()}_Importance'
                    raw_self_importance = pd.DataFrame({
                        'Feature': selected_features,
                        imp_name: self_imp
                    }).sort_values(imp_name, ascending=False)
                    imp_col = imp_name
                else:
                    print( f"Warning: importance length ({len(self_imp)}) does not match the number of features ({len(selected_features)}); skipping model-specific importance.")
            elif classifier_type == 'mlp' and hasattr(raw_model, 'coefs_'):
                first_layer_weights = np.abs(raw_model.coefs_[0])
                self_imp = np.sum(first_layer_weights, axis=0)
                if len(self_imp) == len(selected_features):
                    raw_self_importance = pd.DataFrame({
                        'Feature': selected_features,
                        'MLP_Weight_Importance': self_imp
                    }).sort_values('MLP_Weight_Importance', ascending=False)
                    imp_col = 'MLP_Weight_Importance'
                else:
                    print(f"Warning: MLP weight length ({len(self_imp)}) does not match the number of features ({len(selected_features)}); skipping model-specific importance.")
            elif classifier_type == 'svm' and hasattr(raw_model, 'coef_') and raw_model.kernel == 'linear':
                self_imp = np.abs(raw_model.coef_[0])
                if len(self_imp) == len(selected_features):
                    raw_self_importance = pd.DataFrame({
                        'Feature': selected_features,
                        'SVM_Coef_Importance': self_imp
                    }).sort_values('SVM_Coef_Importance', ascending=False)
                    imp_col = 'SVM_Coef_Importance'
                else:
                    print(f"Warning: SVM coefficient length ({len(self_imp)}) does not match the number of features ({len(selected_features)}); skipping model-specific importance.")
            else:
                pass

            if raw_self_importance is not None:
                self_path = os.path.join(fa_dir, f'{classifier_type}_raw_self_importance.csv')
                raw_self_importance.to_csv(self_path, index=False)
                metric_self_dir = os.path.join(metric_dir, "feature_importance")
                metric_self_path = os.path.join(metric_self_dir, f"{model_key}_self_importance.csv")
                raw_self_importance.to_csv(metric_self_path, index=False)
                

                plt.figure(figsize=(12, 8))
                top20 = raw_self_importance.head(20)
                plt.barh(top20['Feature'], top20.iloc[:, 1])
                plt.xlabel(imp_col)
                plt.title(f'Raw Feature Importance ({classifier_type.upper()} Self)')
                plt.gca().invert_yaxis()
                plt.tight_layout()
                img_path = os.path.join(metric_self_dir, f"{model_key}_self_importance.png")
                plt.savefig(img_path, dpi=300, bbox_inches='tight')
                plt.close()
        except Exception as e:
            print(f"Failed to calculate model-specific importance: {str(e)}")
            traceback.print_exc()

    pca_imp_dict = pca_analyzer.analyze_feature_importance()
    pca_composite = pca_imp_dict['composite']
    if hasattr(pca_analyzer, 'original_feature_names') and pca_analyzer.original_feature_names is not None:
        pca_composite_dict = dict(zip(pca_analyzer.original_feature_names, pca_composite))
        pca_composite_selected = [pca_composite_dict.get(feat, 0.0) for feat in selected_features]
    else:
        if len(pca_composite) >= len(selected_features):
            pca_composite_selected = pca_composite[:len(selected_features)]
        else:
            pca_composite_selected = list(pca_composite) + [0.0] * (len(selected_features) - len(pca_composite))

    pca_orig_df = pd.DataFrame({'Feature': selected_features, 'PCA_Composite_Importance': pca_composite_selected})
    pca_orig_df.to_csv(os.path.join(fa_dir, 'pca_original_feature_importance.csv'), index=False)
    pca_metric_path = os.path.join(metric_dir, "feature_importance", f"{model_key}_pca_importance.csv")
    os.makedirs(os.path.dirname(pca_metric_path), exist_ok=True)
    pca_orig_df.to_csv(pca_metric_path, index=False)
    

    if classifier_type in ['rf', 'dt', 'lgb', 'xgb'] and hasattr(model, 'feature_importances_'):
        clf_imp = model.feature_importances_
        if len(clf_imp) == len(pca_cols):
            clf_df = pd.DataFrame({'Principal_Component': pca_cols, f'{classifier_type.upper()}_Importance': clf_imp})
            clf_df.to_csv(os.path.join(fa_dir, f'{classifier_type}_pca_importance.csv'), index=False)
    elif classifier_type == 'mlp' and hasattr(model, 'coefs_') and model.coefs_:
        first_layer_weights = np.abs(model.coefs_[0])
        if first_layer_weights.shape[1] == len(pca_cols):
            mlp_imp = np.sum(first_layer_weights, axis=0)
            mlp_df = pd.DataFrame({'Principal_Component': pca_cols, 'MLP_Importance': mlp_imp})
            mlp_df.to_csv(os.path.join(fa_dir, 'mlp_pca_importance.csv'), index=False)

    corr_list = []
    for i, feat in enumerate(selected_features):
        if len(np.unique(X_train_original[:, i])) > 1:
            corr = np.corrcoef(X_train_original[:, i], y_train_original)[0, 1]
            if not np.isnan(corr):
                corr_list.append({'Feature': feat, 'Correlation_with_Target': corr, 'Absolute_Correlation': abs(corr)})
    corr_df = pd.DataFrame(corr_list).sort_values('Absolute_Correlation', ascending=False)
    corr_df.to_csv(os.path.join(fa_dir, 'feature_correlation.csv'), index=False)

    combined = pd.DataFrame({'Feature': selected_features})
    combined['PCA_Composite_Importance'] = pca_composite_selected
    corr_map = dict(zip(corr_df['Feature'], corr_df['Absolute_Correlation']))
    combined['Absolute_Correlation'] = combined['Feature'].map(corr_map).fillna(0.0)
    for col in ['PCA_Composite_Importance', 'Absolute_Correlation']:
        minv, maxv = combined[col].min(), combined[col].max()
        if maxv > minv:
            combined[f'{col}_norm'] = (combined[col] - minv) / (maxv - minv)
        else:
            combined[f'{col}_norm'] = 0.0
    combined['Composite_Score'] = (combined['PCA_Composite_Importance_norm'] + combined[
        'Absolute_Correlation_norm']) / 2
    combined = combined.sort_values('Composite_Score', ascending=False)
    combined.to_csv(os.path.join(fa_dir, 'combined_feature_importance.csv'), index=False)
    combined_metric_path = os.path.join(metric_dir, "feature_importance", f"{model_key}_combined_importance.csv")
    combined.to_csv(combined_metric_path, index=False)
    

    
    
    max_samples_for_shap = 1000
    if X_train_original.shape[0] > max_samples_for_shap:
        
        indices = np.random.choice(X_train_original.shape[0], max_samples_for_shap, replace=False)
        X_shap = X_train_original[indices]
        y_shap = y_train_original[indices] if y_train_original is not None else None
    else:
        X_shap = X_train_original
        y_shap = y_train_original

    shap_values = None
    explainer = None
    try:

        if classifier_type in ['rf', 'dt', 'lgb', 'xgb']:
            
            explainer = shap.TreeExplainer(raw_model, feature_names=selected_features)
            shap_values = explainer.shap_values(X_shap)
            if isinstance(shap_values, list) and len(shap_values) == 2:
                shap_values = shap_values[1]
            elif isinstance(shap_values, np.ndarray) and shap_values.ndim == 3:
                shap_values = shap_values[:, :, 1]
        elif classifier_type in ['mlp', 'svm']:
            
            background_size = min(100, X_shap.shape[0])
            background_indices = np.random.choice(X_shap.shape[0], background_size, replace=False)
            background_data = X_shap[background_indices]
            explainer = shap.KernelExplainer(raw_model.predict_proba, background_data, feature_names=selected_features)
            shap_values = explainer.shap_values(X_shap)
            if isinstance(shap_values, list) and len(shap_values) == 2:
                shap_values = shap_values[1]
            elif isinstance(shap_values, np.ndarray) and shap_values.ndim == 3:
                shap_values = shap_values[:, :, 1]

        if shap_values is not None:
            if isinstance(shap_values, np.ndarray) and shap_values.ndim == 3:
                shap_values = shap_values[:, :, 1]
            shap_df = pd.DataFrame(shap_values, columns=selected_features)
            shap_df.to_csv(os.path.join(fa_dir, f'{classifier_type}_shap_values.csv'), index=False)
            shap_importance = pd.DataFrame({
                'Feature': selected_features,
                'SHAP_Importance_MeanAbs': np.abs(shap_values).mean(axis=0)
            }).sort_values('SHAP_Importance_MeanAbs', ascending=False)
            shap_importance.to_csv(os.path.join(fa_dir, f'{classifier_type}_shap_importance.csv'), index=False)
            plt.figure(figsize=(12, 8))
            shap.summary_plot(shap_values, X_shap, feature_names=selected_features, plot_type="bar", show=False)
            plt.title(f'SHAP Feature Importance - {classifier_type.upper()}')
            plt.tight_layout()
            shap_bar_path = os.path.join(metric_dir, "feature_importance", f"{model_key}_shap_importance_bar.png")
            os.makedirs(os.path.dirname(shap_bar_path), exist_ok=True)
            plt.savefig(shap_bar_path, dpi=300, bbox_inches='tight')
            plt.close()
            plt.figure(figsize=(12, 8))
            shap.summary_plot(shap_values, X_shap, feature_names=selected_features, show=False)
            plt.title(f'SHAP Summary Plot - {classifier_type.upper()}')
            plt.tight_layout()
            shap_beeswarm_path = os.path.join(metric_dir, "feature_importance", f"{model_key}_shap_beeswarm.png")
            plt.savefig(shap_beeswarm_path, dpi=300, bbox_inches='tight')
            plt.close()
            if 'Spore_Index_log_x_Tmax1' in selected_features:
                spore_idx = selected_features.index('Spore_Index_log_x_Tmax1')
                spore_shap_values = shap_values[:, spore_idx]
                plt.figure(figsize=(10, 6))
                plt.hist(spore_shap_values, bins=50, alpha=0.7, color='steelblue', edgecolor='black')
                plt.xlabel('SHAP Value for Spore_Index_log_x_Tmax1')
                plt.ylabel('Frequency')
                plt.title(f'Distribution of SHAP Values for Spore_Index_log_x_Tmax1 - {classifier_type.upper()}')
                plt.grid(True, alpha=0.3)
                plt.tight_layout()
                spore_hist_path = os.path.join(metric_dir, "feature_importance",
                                               f"{model_key}_spore_index_shap_hist.png")
                plt.savefig(spore_hist_path, dpi=300, bbox_inches='tight')
                plt.close()
                plt.figure(figsize=(10, 6))
                shap.dependence_plot('Spore_Index_log_x_Tmax1', shap_values, X_shap, feature_names=selected_features, show=False)
                plt.title(f'SHAP Dependence Plot for Spore_Index_log_x_Tmax1 - {classifier_type.upper()}')
                plt.tight_layout()
                spore_dep_path = os.path.join(metric_dir, "feature_importance",
                                              f"{model_key}_spore_index_dependence.png")
                plt.savefig(spore_dep_path, dpi=300, bbox_inches='tight')
                plt.close()
    except Exception as e:
        print(f"Failed to calculate SHAP values: {str(e)}")
        traceback.print_exc()
        shap_values = None

    report = {
        "analysis_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_features": len(selected_features),
        "analysis_methods": ["PCA_composite_importance", "feature_correlation", "composite_importance", "permutation_importance_original_features",
                     "model_specific_importance", "SHAP_importance"] if raw_perm_importance is not None else ["PCA_composite_importance",
                                                                                                "feature_correlation",
                                                                                                "composite_importance",
                                                                                                "SHAP_importance"]
    }
    if raw_perm_importance is not None:
        report["top10_permutation_importance"] = raw_perm_importance.head(10).to_dict('records')
    if raw_self_importance is not None:
        report["top10_model_specific_importance"] = raw_self_importance.head(10).to_dict('records')
    report["top10_PCA_original_feature_importance"] = pca_orig_df.head(10).to_dict('records')
    report["top10_feature_correlation"] = corr_df.head(10).to_dict('records')
    report["top10_composite_importance"] = combined.head(10)[['Feature', 'Composite_Score']].to_dict('records')
    if shap_values is not None:
        shap_importance = pd.DataFrame({
            'Feature': selected_features,
            'SHAP_Importance_MeanAbs': np.abs(shap_values).mean(axis=0)
        }).sort_values('SHAP_Importance_MeanAbs', ascending=False)
        report["top10_SHAP_importance"] = shap_importance.head(10).to_dict('records')
    with open(os.path.join(fa_dir, 'factor_analysis_report.json'), 'w', encoding='utf-8') as f:
        json.dump(convert_types(report), f, indent=4, ensure_ascii=False)
    md = f"""# Feature Importance Analysis Report

**generated_at**: {report['analysis_time']}

## Analysis overview
- **total_features**: {report['total_features']}
- **analysis_methods**: {', '.join(report['analysis_methods'])}

## Feature importance rankings

### Top 10 permutation importance (original features)
{generate_md_table(report.get('top10_permutation_importance', []))}

### Top 10 model-specific importance
{generate_md_table(report.get('top10_model_specific_importance', []))}

### Top 10 PCA-based original-feature importance
{generate_md_table(report.get('top10_PCA_original_feature_importance', []))}

### Top 10 feature correlation
{generate_md_table(report.get('top10_feature_correlation', []))}

### Top 10 composite importance
{generate_md_table(report.get('top10_composite_importance', []))}

### Top 10 SHAP importance
{generate_md_table(report.get('top10_SHAP_importance', []))}

"""
    with open(os.path.join(fa_dir, 'factor_analysis_report.md'), 'w', encoding='utf-8') as f:
        f.write(md)
    return raw_perm_importance, shap_values, selected_features

def generate_md_table(data):
    if not data:
        return "No data"
    headers = list(data[0].keys())
    table = "| " + " | ".join(headers) + " |\n"
    table += "|" + "|".join(["---"] * len(headers)) + "|\n"
    for row in data:
        table += "| " + " | ".join(str(x) for x in row.values()) + " |\n"
    return table

def evaluate_final_model_on_test_set(model, pca_analyzer, selector, core_indices,
                                     X_test, y_test, result_dir, classifier_type, metric_dir, model_key, threshold=0.5):
    
    
    
    if X_test is None or y_test is None:
        print(" The test set is empty; evaluation cannot be performed.")
        return None, None

    try:
        X_test_core = X_test[:, core_indices]
        if hasattr(selector, 'selected_indices'):
            X_test_selected = X_test_core[:, selector.selected_indices]
        else:
            X_test_selected = selector.transform(X_test_core)

        X_test_pca = pca_analyzer.transform(X_test_selected)
        test_probs_full = model.predict_proba(X_test_pca)
        test_probs = test_probs_full[:, 1]
        test_preds = (test_probs >= threshold).astype(int)
        test_metrics = compute_binary_metrics(y_test, test_preds, test_probs)
        test_acc = test_metrics['Accuracy']
        test_auc = test_metrics['ROC-AUC']
        print(
            f" Test-set metrics: Accuracy={test_acc:.4f}, "
            f"BA={test_metrics['BA']:.4f}, G-mean={test_metrics['G-mean']:.4f}, "
            f"F1={test_metrics['F1-score']:.4f}, Recall={test_metrics['Recall']:.4f}, "
            f"AUC={test_auc:.4f}"
        )

        os.makedirs(os.path.join(result_dir, "test_results"), exist_ok=True)
        with open(os.path.join(result_dir, "test_results", "final_model_test_results.json"), 'w', encoding='utf-8') as f:
            json.dump(convert_types({
                **test_metrics,
                'test_accuracy': float(test_acc),
                'test_auc': float(test_auc),
                'test_size': int(X_test.shape[0]),
                'threshold': threshold,
                'evaluation_dataset': 'independent_original_test_set'
            }), f, indent=4, ensure_ascii=False)

        roc_model_path = os.path.join(result_dir, "roc_curves", "final_model_test_roc.png")
        if plot_roc_curve_binary(y_test, test_probs_full, model.classes_,
                                 'ROC Curve - Final Model (Test)', roc_model_path):
            roc_metric_path = os.path.join(metric_dir, "roc_curves", f"{model_key}_test_roc.png")
            if os.path.exists(roc_model_path):
                shutil.copy2(roc_model_path, roc_metric_path)

                roc_data_src = roc_model_path.replace('.png', '_data.csv')
                roc_data_dst = roc_metric_path.replace('.png', '_data.csv')
                if os.path.exists(roc_data_src):
                    shutil.copy2(roc_data_src, roc_data_dst)
        else:
            print(f" Failed to plot the test-set ROC curve; skipping copy")

        pr_dir = os.path.join(result_dir, "pr_curves")
        os.makedirs(pr_dir, exist_ok=True)
        pr_model_path = os.path.join(pr_dir, "final_model_test_pr.png")
        if plot_pr_curve_binary(y_test, test_probs_full, model.classes_,
                                'PR Curve - Final Model (Test)', pr_model_path):

            pr_metric_dir = os.path.join(metric_dir, "pr_curves")
            os.makedirs(pr_metric_dir, exist_ok=True)
            pr_metric_path = os.path.join(pr_metric_dir, f"{model_key}_test_pr.png")
            shutil.copy2(pr_model_path, pr_metric_path)

            pr_data_src = pr_model_path.replace('.png', '_data.csv')
            pr_data_dst = pr_metric_path.replace('.png', '_data.csv')
            if os.path.exists(pr_data_src):
                shutil.copy2(pr_data_src, pr_data_dst)
        else:
            print(f" Failed to plot the test-set PR curve; skipping copy")

        unique_cls = np.unique(y_test)
        class_labels = [str(c) for c in unique_cls]
        save_confusion_matrix(y_test, test_preds,
                              os.path.join(result_dir, "confusion_matrices", "final_model_test_cm.png"),
                              classes=class_labels, title="Final Model - Test CM")
        save_confusion_matrix(y_test, test_preds,
                              os.path.join(result_dir, "confusion_matrices", "final_model_test_cm_norm.png"),
                              classes=class_labels, normalize=True, title="Final Model - Test CM (Norm)")
        cm_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_test_cm.png")
        save_confusion_matrix(y_test, test_preds, cm_path, classes=class_labels, title="Final Model - Test CM")
        cm_norm_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_test_cm_norm.png")
        save_confusion_matrix(y_test, test_preds, cm_norm_path, classes=class_labels, normalize=True,
                              title="Final Model - Test CM (Norm)")

        excel_path = os.path.join(result_dir, "confusion_matrices", "final_model_test_cm.xlsx")
        save_confusion_matrix_excel(y_test, test_preds, excel_path, classes=class_labels, normalize=False)
        excel_norm_path = os.path.join(result_dir, "confusion_matrices", "final_model_test_cm_norm.xlsx")
        save_confusion_matrix_excel(y_test, test_preds, excel_norm_path, classes=class_labels, normalize=True)
        metric_excel_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_test_cm.xlsx")
        save_confusion_matrix_excel(y_test, test_preds, metric_excel_path, classes=class_labels, normalize=False)
        metric_excel_norm_path = os.path.join(metric_dir, "confusion_matrices", f"{model_key}_test_cm_norm.xlsx")
        save_confusion_matrix_excel(y_test, test_preds, metric_excel_norm_path, classes=class_labels, normalize=True)

        return test_acc, test_auc
    except Exception as e:
        print(f" Test-set evaluation failed: {str(e)}")
        traceback.print_exc()
        return None, None

def collect_model_detailed_metrics(model_path, model_name, output_dir):
    model_output_dir = os.path.join(output_dir, model_name)
    os.makedirs(model_output_dir, exist_ok=True)

    kfold_csv = os.path.join(model_path, 'results', 'kfold_cv_results.csv')
    if os.path.exists(kfold_csv):
        shutil.copy2(kfold_csv, os.path.join(model_output_dir, 'kfold_cv_results.csv'))

    kfold_json = os.path.join(model_path, 'results', 'kfold_results.json')
    if os.path.exists(kfold_json):
        shutil.copy2(kfold_json, os.path.join(model_output_dir, 'kfold_results.json'))

    cm_dir = os.path.join(model_path, 'confusion_matrices')
    if os.path.exists(cm_dir):
        for file in os.listdir(cm_dir):
            if file.endswith('.xlsx'):
                shutil.copy2(os.path.join(cm_dir, file), os.path.join(model_output_dir, file))

    roc_dir = os.path.join(model_path, 'roc_curves')
    if os.path.exists(roc_dir):
        for file in os.listdir(roc_dir):
            if file.endswith('_data.csv'):
                shutil.copy2(os.path.join(roc_dir, file), os.path.join(model_output_dir, file))

    pr_dir = os.path.join(model_path, 'pr_curves')
    if os.path.exists(pr_dir):
        pr_data_dir = os.path.join(model_output_dir, 'pr_curves')
        os.makedirs(pr_data_dir, exist_ok=True)
        for file in os.listdir(pr_dir):
            if file.endswith('.csv') or file.endswith('.png'):
                shutil.copy2(os.path.join(pr_dir, file), os.path.join(pr_data_dir, file))

    if os.path.exists(kfold_json):
        with open(kfold_json, 'r') as f:
            data = json.load(f)
        fold_metrics = []
        for item in data:
            fold_metrics.append({
                'Fold': item['fold'],
                'Accuracy': item['accuracy'],
                'AUC': item.get('auc', None)
            })
        df_fold = pd.DataFrame(fold_metrics)
        df_fold.to_csv(os.path.join(model_output_dir, 'fold_metrics.csv'), index=False)

    

def generate_all_models_comparison(result_paths, base_config, output_dir, all_raw_importances_dict,
                                   all_shap_importances_dict):
    
    
    

    comparison_data = []
    for path in result_paths:
        dir_name = os.path.basename(path)
        if '+' in dir_name:
            fs, clf = dir_name.split('+')
        else:
            continue

        summary_json = os.path.join(path, "summary", "final_report.json")
        test_json = os.path.join(path, "test_results", "final_model_test_results.json")
        model_info_json = os.path.join(path, "models", "final_model_info.json")

        cv_acc = None
        train_acc = None
        test_acc = None
        test_auc = None
        num_features = None

        if os.path.exists(summary_json):
            try:
                with open(summary_json, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    cv_acc = float(data.get("model_performance", {}).get("mean_cross_validation_accuracy", "N/A").replace("N/A", "nan"))
                    train_acc = float(data.get("model_performance", {}).get("training_accuracy", "N/A").replace("N/A", "nan"))
                    test_acc = float(data.get("model_performance", {}).get("test_accuracy", "N/A").replace("N/A", "nan"))
                    test_auc = float(data.get("model_performance", {}).get("test_auc", "N/A").replace("N/A", "nan"))
            except Exception as e:
                print(f"Failed to read {summary_json}: {e}")

        if os.path.exists(test_json):
            try:
                with open(test_json, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    test_acc = data.get("test_accuracy", test_acc)
                    test_auc = data.get("test_auc", test_auc)
            except Exception as e:
                print(f"Failed to read {test_json}: {e}")

        if os.path.exists(model_info_json):
            try:
                with open(model_info_json, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    num_features = data.get("num_selected_features", None)
            except Exception as e:
                print(f"Failed to read {model_info_json}: {e}")

        comparison_data.append({
            "FeatureSelector": fs.upper(),
            "Classifier": clf.upper(),
            "CV_Accuracy": cv_acc,
            "Train_Accuracy": train_acc,
            "Test_Accuracy": test_acc,
            "Test_AUC": test_auc,
            "Num_Features": num_features,
            "ResultDir": path
        })

    if not comparison_data:
        print(" No valid model-performance data found; skipping the comparison summary.")
        return

    df = pd.DataFrame(comparison_data)
    numeric_cols = ['CV_Accuracy', 'Train_Accuracy', 'Test_Accuracy', 'Test_AUC']
    for col in numeric_cols:
        df[col] = pd.to_numeric(df[col], errors='coerce')
    df = df.dropna(subset=numeric_cols, how='any')
    if df.empty:
        print(" All model-performance data are missing; comparison plots cannot be generated.")
        return

    df = df.sort_values(["FeatureSelector", "Classifier"])
    os.makedirs(output_dir, exist_ok=True)
    df.to_csv(os.path.join(output_dir, "all_models_performance.csv"), index=False, encoding='utf-8-sig')

    plt.figure(figsize=(12, 6))
    x = np.arange(len(df))
    width = 0.35
    plt.bar(x - width / 2, df['Test_Accuracy'], width, label='Test Accuracy', color='steelblue')
    plt.bar(x + width / 2, df['Test_AUC'], width, label='Test AUC', color='coral')
    plt.xticks(x, df['FeatureSelector'] + '+' + df['Classifier'], rotation=45, ha='right')
    plt.ylabel('Score')
    plt.title('Comparison of Test Accuracy and AUC across Models')
    plt.legend()
    plt.grid(axis='y', alpha=0.3)
    plt.tight_layout()
    plt.savefig(os.path.join(output_dir, "test_performance_bar.png"), dpi=300)
    plt.close()

    pivot_acc = df.pivot(index='FeatureSelector', columns='Classifier', values='Test_Accuracy')
    plt.figure(figsize=(8, 6))
    sns.heatmap(pivot_acc, annot=True, fmt='.4f', cmap='YlGnBu', cbar_kws={'label': 'Test Accuracy'})
    plt.title('Test Accuracy Heatmap')
    plt.tight_layout()
    plt.savefig(os.path.join(output_dir, "test_accuracy_heatmap.png"), dpi=300)
    plt.close()

    if df['Num_Features'].notna().any():
        plt.figure(figsize=(12, 6))
        plt.bar(df['FeatureSelector'] + '+' + df['Classifier'], df['Num_Features'], color='green', alpha=0.7)
        plt.ylabel('Number of Selected Features')
        plt.title('Number of Features after Selection across Models')
        plt.xticks(rotation=45, ha='right')
        plt.grid(axis='y', alpha=0.3)
        plt.tight_layout()
        plt.savefig(os.path.join(output_dir, "feature_counts.png"), dpi=300)
        plt.close()

    try:
        plt.figure(figsize=(10, 8))
        colors = plt.cm.tab20(np.linspace(0, 1, len(result_paths)))
        for i, path in enumerate(result_paths):
            dir_name = os.path.basename(path)
            roc_data_file = os.path.join(path, "roc_curves", "final_model_test_roc_data.csv")
            if os.path.exists(roc_data_file):
                roc_data = pd.read_csv(roc_data_file, encoding='utf-8', comment='#')
                roc_data['False_Positive_Rate'] = pd.to_numeric(roc_data['False_Positive_Rate'], errors='coerce')
                roc_data['True_Positive_Rate'] = pd.to_numeric(roc_data['True_Positive_Rate'], errors='coerce')
                roc_data = roc_data.dropna(subset=['False_Positive_Rate', 'True_Positive_Rate'])

                auc_val = None
                with open(roc_data_file, 'r', encoding='utf-8') as f:
                    lines = f.readlines()
                    for line in lines:
                        if line.startswith('AUC,'):
                            auc_val = float(line.strip().split(',')[1])
                            break
                if auc_val is None:
                    auc_val = 0.0
                plt.plot(roc_data['False_Positive_Rate'], roc_data['True_Positive_Rate'],
                         color=colors[i], lw=2, label=f'{dir_name} (AUC={auc_val:.3f})')
        plt.plot([0, 1], [0, 1], 'k--', lw=2, label='Random')
        plt.xlim([0.0, 1.0])
        plt.ylim([0.0, 1.05])
        plt.xlabel('False Positive Rate', fontsize=12, fontfamily='Arial')
        plt.ylabel('True Positive Rate', fontsize=12, fontfamily='Arial')
        plt.title('Combined ROC Curves (Test Set)', fontsize=14, fontweight='bold', fontfamily='Arial')
        plt.legend(loc='lower right', prop={'family': 'Arial', 'size': 10})
        plt.grid(True, alpha=0.3)
        plt.tight_layout()
        combined_roc_path = os.path.join(output_dir, "combined_roc.png")
        plt.savefig(combined_roc_path, dpi=300, bbox_inches='tight')
        plt.close()
        
    except Exception as e:
        print(f"Failed to plot the combined ROC curve: {str(e)}")

    try:
        plt.figure(figsize=(10, 8))
        colors = plt.cm.tab20(np.linspace(0, 1, len(result_paths)))
        for i, path in enumerate(result_paths):
            dir_name = os.path.basename(path)
            pr_data_file = os.path.join(path, "pr_curves", "final_model_test_pr_data.csv")
            if os.path.exists(pr_data_file):
                pr_data = pd.read_csv(pr_data_file, encoding='utf-8', comment='#')
                pr_data['Recall'] = pd.to_numeric(pr_data['Recall'], errors='coerce')
                pr_data['Precision'] = pd.to_numeric(pr_data['Precision'], errors='coerce')
                pr_data = pr_data.dropna(subset=['Recall', 'Precision'])

                ap_val = None
                with open(pr_data_file, 'r', encoding='utf-8') as f:
                    lines = f.readlines()
                    for line in lines:
                        if line.startswith('Average_Precision,'):
                            ap_val = float(line.strip().split(',')[1])
                            break
                if ap_val is None:
                    ap_val = 0.0
                plt.plot(pr_data['Recall'], pr_data['Precision'],
                         color=colors[i], lw=2, label=f'{dir_name} (AP={ap_val:.3f})')
        plt.xlim([0.0, 1.0])
        plt.ylim([0.0, 1.05])
        plt.xlabel('Recall', fontsize=12, fontfamily='Arial')
        plt.ylabel('Precision', fontsize=12, fontfamily='Arial')
        plt.title('Combined Precision-Recall Curves (Test Set)', fontsize=14, fontweight='bold', fontfamily='Arial')
        plt.legend(loc='lower left', prop={'family': 'Arial', 'size': 10})
        plt.grid(True, alpha=0.3)
        plt.tight_layout()
        combined_pr_path = os.path.join(output_dir, "combined_pr.png")
        plt.savefig(combined_pr_path, dpi=300, bbox_inches='tight')
        plt.close()
        
    except Exception as e:
        print(f"Failed to plot the combined PR curve: {str(e)}")

    if all_raw_importances_dict:
        all_features = set()
        for imp_df in all_raw_importances_dict.values():
            if imp_df is not None:
                all_features.update(imp_df['Feature'].tolist())
        all_features = sorted(list(all_features))

        heatmap_data = {}
        for model_key, imp_df in all_raw_importances_dict.items():
            if imp_df is None:
                continue
            imp_map = dict(zip(imp_df['Feature'], imp_df['Permutation_Importance_Mean']))
            heatmap_data[model_key] = [imp_map.get(f, 0.0) for f in all_features]

        if heatmap_data:
            heatmap_df = pd.DataFrame(heatmap_data, index=all_features)
            heatmap_df_norm = (heatmap_df - heatmap_df.min()) / (heatmap_df.max() - heatmap_df.min() + 1e-8)
            plt.figure(figsize=(max(10, len(heatmap_data) * 0.8), max(10, len(all_features) * 0.2)))
            sns.heatmap(heatmap_df_norm, annot=False, cmap='viridis',
                        cbar_kws={'label': 'Normalized Permutation Importance'})
            plt.title('Raw Feature Permutation Importance across Models')
            plt.xlabel('Model Combination')
            plt.ylabel('Original Features')
            plt.xticks(rotation=45, ha='right')
            plt.tight_layout()
            heatmap_path = os.path.join(output_dir, "raw_feature_importance_heatmap.png")
            plt.savefig(heatmap_path, dpi=300, bbox_inches='tight')
            plt.close()
            heatmap_df.to_csv(os.path.join(output_dir, "raw_feature_importance_matrix.csv"))
            

    if all_shap_importances_dict:
        all_features = set()
        for shap_data in all_shap_importances_dict.values():
            if shap_data is not None and 'shap_values' in shap_data and shap_data['shap_values'] is not None:
                all_features.update(shap_data['feature_names'])
        all_features = sorted(list(all_features))

        shap_heatmap_data = {}
        for model_key, shap_data in all_shap_importances_dict.items():
            if shap_data is None or 'shap_values' not in shap_data or shap_data['shap_values'] is None:
                continue
            shap_values = shap_data['shap_values']
            feature_names = shap_data['feature_names']
            shap_importance = np.abs(shap_values).mean(axis=0)
            shap_map = dict(zip(feature_names, shap_importance))
            shap_heatmap_data[model_key] = [shap_map.get(f, 0.0) for f in all_features]

        if shap_heatmap_data:
            shap_heatmap_df = pd.DataFrame(shap_heatmap_data, index=all_features)
            shap_heatmap_df_norm = (shap_heatmap_df - shap_heatmap_df.min()) / (shap_heatmap_df.max() - shap_heatmap_df.min() + 1e-8)
            plt.figure(figsize=(max(10, len(shap_heatmap_data) * 0.8), max(10, len(all_features) * 0.2)))
            sns.heatmap(shap_heatmap_df_norm, annot=False, cmap='coolwarm',
                        cbar_kws={'label': 'Normalized SHAP Importance'})
            plt.title('SHAP Feature Importance across Models')
            plt.xlabel('Model Combination')
            plt.ylabel('Original Features')
            plt.xticks(rotation=45, ha='right')
            plt.tight_layout()
            shap_heatmap_path = os.path.join(output_dir, "shap_feature_importance_heatmap.png")
            plt.savefig(shap_heatmap_path, dpi=300, bbox_inches='tight')
            plt.close()
            shap_heatmap_df.to_csv(os.path.join(output_dir, "shap_feature_importance_matrix.csv"))
            

    spore_data = []
    if all_raw_importances_dict:
        for model_key, imp_df in all_raw_importances_dict.items():
            if imp_df is None:
                continue
            spore_row = imp_df[imp_df['Feature'] == 'Spore_Index_log_x_Tmax1']
            if not spore_row.empty:
                importance = spore_row.iloc[0]['Permutation_Importance_Mean']
                std = spore_row.iloc[0]['Permutation_Importance_Std']
                imp_df_sorted = imp_df.sort_values('Permutation_Importance_Mean', ascending=False)
                rank_series = imp_df_sorted.reset_index(drop=True).reset_index()
                rank_series.columns = ['index', 'Feature'] + list(imp_df_sorted.columns[1:])
                rank_row = rank_series[rank_series['Feature'] == 'Spore_Index_log_x_Tmax1']
                if not rank_row.empty:
                    rank = rank_row.iloc[0]['index'] + 1
                else:
                    rank = None
                total_features = len(imp_df)
                percentile = (total_features - rank) / total_features * 100 if rank is not None else 0.0
            else:
                importance, std, rank, percentile = 0.0, 0.0, None, 0.0
                total_features = len(imp_df) if imp_df is not None else 0
            spore_data.append({
                'Model': model_key,
                'Importance': importance,
                'Std': std,
                'Rank': rank,
                'Percentile': percentile,
                'Total_Features': total_features
            })
        if spore_data:
            spore_df = pd.DataFrame(spore_data)
            spore_df.to_csv(os.path.join(output_dir, 'Spore_Index_permutation_importance.csv'), index=False)
            plt.figure(figsize=(12, 6))
            models = spore_df['Model']
            imp_vals = spore_df['Importance']
            std_vals = spore_df['Std']
            plt.bar(models, imp_vals, yerr=std_vals, capsize=5, color='steelblue', alpha=0.7)
            plt.xlabel('Model Combination')
            plt.ylabel('Permutation Importance (AUC drop)')
            plt.title('Spore_Index_log_x_Tmax1 Permutation Importance Across Models')
            plt.xticks(rotation=45, ha='right')
            plt.grid(axis='y', alpha=0.3)
            for i, (model, imp, rank) in enumerate(zip(models, imp_vals, spore_df['Rank'])):
                if rank is not None and not np.isnan(rank):
                    plt.text(i, imp + 0.01, f'Rank {int(rank)}', ha='center', fontsize=9)
            plt.tight_layout()
            plt.savefig(os.path.join(output_dir, 'Spore_Index_permutation_importance_bar.png'), dpi=300)
            plt.close()
            

    self_spore_data = []
    for path in result_paths:
        dir_name = os.path.basename(path)
        self_file = os.path.join(path, "factor_analysis", f"{dir_name.split('+')[1]}_raw_self_importance.csv")
        if not os.path.exists(self_file):
            self_file = os.path.join(output_dir, "feature_importance", f"{dir_name}_self_importance.csv")
        if os.path.exists(self_file):
            df_self = pd.read_csv(self_file)
            spore_row = df_self[df_self['Feature'] == 'Spore_Index_log_x_Tmax1']
            if not spore_row.empty:
                imp_col = [c for c in df_self.columns if c != 'Feature'][0]
                importance = spore_row.iloc[0][imp_col]
                df_sorted = df_self.sort_values(imp_col, ascending=False)
                rank_series = df_sorted.reset_index(drop=True).reset_index()
                rank_series.columns = ['index', 'Feature'] + list(df_sorted.columns[1:])
                rank_row = rank_series[rank_series['Feature'] == 'Spore_Index_log_x_Tmax1']
                if not rank_row.empty:
                    rank = rank_row.iloc[0]['index'] + 1
                else:
                    rank = None
                total_features = len(df_self)
                self_spore_data.append({
                    'Model': dir_name,
                    'Importance_Type': imp_col,
                    'Importance': importance,
                    'Rank': rank,
                    'Total_Features': total_features
                })
    if self_spore_data:
        self_spore_df = pd.DataFrame(self_spore_data)
        self_spore_df.to_csv(os.path.join(output_dir, 'Spore_Index_self_importance.csv'), index=False)
        plt.figure(figsize=(12, 6))
        models = self_spore_df['Model']
        imp_vals = self_spore_df['Importance']
        plt.bar(models, imp_vals, color='coral', alpha=0.7)
        plt.xlabel('Model Combination')
        plt.ylabel('Self Importance Score')
        plt.title('Spore_Index_log_x_Tmax1 Self Importance Across Models')
        plt.xticks(rotation=45, ha='right')
        plt.grid(axis='y', alpha=0.3)
        for i, (model, imp, rank) in enumerate(zip(models, imp_vals, self_spore_df['Rank'])):
            if rank is not None and not np.isnan(rank):
                plt.text(i, imp + 0.01, f'Rank {int(rank)}', ha='center', fontsize=9)
        plt.tight_layout()
        plt.savefig(os.path.join(output_dir, 'Spore_Index_self_importance_bar.png'), dpi=300)
        plt.close()
        

    spore_shap_data = []
    if all_shap_importances_dict:
        for model_key, shap_data in all_shap_importances_dict.items():
            if shap_data is None or 'shap_values' not in shap_data or shap_data['shap_values'] is None:
                continue
            shap_values = shap_data['shap_values']
            feature_names = shap_data['feature_names']
            if 'Spore_Index_log_x_Tmax1' in feature_names:
                spore_idx = feature_names.index('Spore_Index_log_x_Tmax1')
                spore_shap = shap_values[:, spore_idx]
                spore_shap_mean = np.mean(spore_shap)
                spore_shap_abs_mean = np.mean(np.abs(spore_shap))
                spore_shap_std = np.std(spore_shap)
                shap_importance = np.abs(shap_values).mean(axis=0)
                shap_importance_df = pd.DataFrame({
                    'Feature': feature_names,
                    'SHAP_Importance': shap_importance
                }).sort_values('SHAP_Importance', ascending=False)
                rank_series = shap_importance_df.reset_index(drop=True).reset_index()
                rank_series.columns = ['index', 'Feature', 'SHAP_Importance']
                rank_row = rank_series[rank_series['Feature'] == 'Spore_Index_log_x_Tmax1']
                if not rank_row.empty:
                    rank = rank_row.iloc[0]['index'] + 1
                else:
                    rank = None
                total_features = len(feature_names)
                spore_shap_data.append({
                    'Model': model_key,
                    'SHAP_Mean': spore_shap_mean,
                    'SHAP_Abs_Mean': spore_shap_abs_mean,
                    'SHAP_Std': spore_shap_std,
                    'Rank': rank,
                    'Total_Features': total_features
                })
    if spore_shap_data:
        spore_shap_df = pd.DataFrame(spore_shap_data)
        spore_shap_df.to_csv(os.path.join(output_dir, 'Spore_Index_shap_importance.csv'), index=False)
        plt.figure(figsize=(12, 6))
        models = spore_shap_df['Model']
        imp_vals = spore_shap_df['SHAP_Abs_Mean']
        plt.bar(models, imp_vals, color='purple', alpha=0.7)
        plt.xlabel('Model Combination')
        plt.ylabel('SHAP Absolute Mean Value')
        plt.title('Spore_Index_log_x_Tmax1 SHAP Importance Across Models')
        plt.xticks(rotation=45, ha='right')
        plt.grid(axis='y', alpha=0.3)
        for i, (model, imp, rank) in enumerate(zip(models, imp_vals, spore_shap_df['Rank'])):
            if rank is not None and not np.isnan(rank):
                plt.text(i, imp + 0.01, f'Rank {int(rank)}', ha='center', fontsize=9)
        plt.tight_layout()
        plt.savefig(os.path.join(output_dir, 'Spore_Index_shap_importance_bar.png'), dpi=300)
        plt.close()
        

    detailed_metrics_dir = os.path.join(output_dir, 'detailed_metrics')
    os.makedirs(detailed_metrics_dir, exist_ok=True)
    for path in result_paths:
        dir_name = os.path.basename(path)
        collect_model_detailed_metrics(path, dir_name, detailed_metrics_dir)

    
    comprehensive_excel_path = os.path.join(output_dir, 'comprehensive_metrics.xlsx')
    with pd.ExcelWriter(comprehensive_excel_path, engine='openpyxl') as writer:
        df_perf = df[['FeatureSelector', 'Classifier', 'CV_Accuracy', 'Train_Accuracy', 'Test_Accuracy', 'Test_AUC', 'Num_Features']].copy()
        df_perf.to_excel(writer, sheet_name='Model_Performance', index=False)

        fold_all = []
        for path in result_paths:
            model_name = os.path.basename(path)
            kfold_json = os.path.join(path, 'results', 'kfold_results.json')
            if os.path.exists(kfold_json):
                with open(kfold_json, 'r') as f:
                    data = json.load(f)
                for item in data:
                    fold_all.append({
                        'Model': model_name,
                        'Fold': item['fold'],
                        'Accuracy': item['accuracy'],
                        'AUC': item.get('auc', None)
                    })
        if fold_all:
            df_fold_all = pd.DataFrame(fold_all)
            df_fold_all.to_excel(writer, sheet_name='Fold_Accuracy_AUC', index=False)

        train_cm_data = []
        for path in result_paths:
            model_name = os.path.basename(path)
            cm_excel = os.path.join(path, 'confusion_matrices', 'final_model_training_cm.xlsx')
            if os.path.exists(cm_excel):
                df_cm = pd.read_excel(cm_excel, sheet_name='Confusion_Matrix', index_col=0)
                for true_label in df_cm.index:
                    if true_label == 'Total':
                        continue
                    for pred_label in df_cm.columns:
                        if pred_label == 'Total':
                            continue
                        train_cm_data.append({
                            'Model': model_name,
                            'True_Label': true_label,
                            'Predicted_Label': pred_label,
                            'Count': df_cm.loc[true_label, pred_label]
                        })
        if train_cm_data:
            df_train_cm = pd.DataFrame(train_cm_data)
            df_train_cm.to_excel(writer, sheet_name='Train_Confusion_Matrices', index=False)

        test_cm_data = []
        for path in result_paths:
            model_name = os.path.basename(path)
            cm_excel = os.path.join(path, 'confusion_matrices', 'final_model_test_cm.xlsx')
            if os.path.exists(cm_excel):
                df_cm = pd.read_excel(cm_excel, sheet_name='Confusion_Matrix', index_col=0)
                for true_label in df_cm.index:
                    if true_label == 'Total':
                        continue
                    for pred_label in df_cm.columns:
                        if pred_label == 'Total':
                            continue
                        test_cm_data.append({
                            'Model': model_name,
                            'True_Label': true_label,
                            'Predicted_Label': pred_label,
                            'Count': df_cm.loc[true_label, pred_label]
                        })
        if test_cm_data:
            df_test_cm = pd.DataFrame(test_cm_data)
            df_test_cm.to_excel(writer, sheet_name='Test_Confusion_Matrices', index=False)

    print(f" Comprehensive evaluation workbook saved to: {comprehensive_excel_path}")

    md_content = f"""# Performance comparison of 12 model combinations

**generated_at**: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

## Performance overview

| Feature selector | Classifier | CV accuracy | Training accuracy | Test accuracy | Test AUC | Final number of features |
|----------|--------|----------------|---------------|---------------|------------|------------|
"""
    for _, row in df.iterrows():
        md_content += f"| {row['FeatureSelector']} | {row['Classifier']} | {row['CV_Accuracy']:.4f} | {row['Train_Accuracy']:.4f} | {row['Test_Accuracy']:.4f} | {row['Test_AUC']:.4f} | {row['Num_Features']} |\n"

    best_acc_idx = df['Test_Accuracy'].idxmax()
    best_acc_model = f"{df.loc[best_acc_idx, 'FeatureSelector']}+{df.loc[best_acc_idx, 'Classifier']}"
    best_acc = df.loc[best_acc_idx, 'Test_Accuracy']
    best_auc_idx = df['Test_AUC'].idxmax()
    best_auc_model = f"{df.loc[best_auc_idx, 'FeatureSelector']}+{df.loc[best_auc_idx, 'Classifier']}"
    best_auc = df.loc[best_auc_idx, 'Test_AUC']

    md_content += f"""
## Best model

- **Highest test accuracy**: {best_acc_model} ({best_acc:.4f})
- **Highest test AUC**: {best_auc_model} ({best_auc:.4f})

## Spore_Index importance analysis

See the following files for detailed results:
- `Spore_Index_permutation_importance.csv` and `Spore_Index_permutation_importance_bar.png`(permutation importance)
- `Spore_Index_self_importance.csv` and `Spore_Index_self_importance_bar.png`(model-specific importance)
- `Spore_Index_shap_importance.csv` and `Spore_Index_shap_importance_bar.png`(SHAP importance)
- `raw_feature_importance_heatmap.png` and `raw_feature_importance_matrix.csv`(permutation-importance heatmap for all features)
- `shap_feature_importance_heatmap.png` and `shap_feature_importance_matrix.csv`(SHAP-importance heatmap for all features)

## Combined ROC and PR curves

Combined curves are saved as `combined_roc.png` and `combined_pr.png`, showing test-set performance across all models.

## Conclusion

Based on test-set performance, the recommended model is **{best_acc_model}**.

See the model-specific subdirectories for detailed results.
"""
    with open(os.path.join(output_dir, "comparison_report.md"), 'w', encoding='utf-8') as f:
        f.write(md_content)

    print(f" Model-comparison summary saved to: {os.path.abspath(output_dir)}")
    return output_dir

def generate_table_s12(result_paths, output_dir):

    

    classifier_order = ['DT', 'LGB', 'MLP', 'RF', 'SVM', 'XGB']
    fs_order = ['RF', 'RFE']
    rows = []

    for path in result_paths:
        model_name = os.path.basename(path)
        if '+' not in model_name:
            continue

        fs, clf = model_name.split('+', 1)
        fs_u, clf_u = fs.upper(), clf.upper()

        train_json = os.path.join(path, 'final_model_details', 'final_train_metrics.json')
        test_json = os.path.join(path, 'test_results', 'final_model_test_results.json')
        model_info_json = os.path.join(path, 'models', 'final_model_info.json')
        kfold_csv = os.path.join(path, 'results', 'kfold_cv_results.csv')

        if not os.path.exists(train_json):
            print(f" {model_name}: Missing {train_json}; skipping Table S12 aggregation")
            continue
        if not os.path.exists(test_json):
            print(f" {model_name}: Missing {test_json}; skipping Table S12 aggregation")
            continue

        with open(train_json, 'r', encoding='utf-8') as f:
            tr = json.load(f)
        with open(test_json, 'r', encoding='utf-8') as f:
            te = json.load(f)

        features = None
        if os.path.exists(model_info_json):
            with open(model_info_json, 'r', encoding='utf-8') as f:
                info = json.load(f)
            features = info.get('num_selected_features', None)

        cv_acc = np.nan
        if os.path.exists(kfold_csv):
            cv_df = pd.read_csv(kfold_csv, encoding='utf-8')
            if 'Accuracy' in cv_df.columns:
                cv_acc = float(pd.to_numeric(cv_df['Accuracy'], errors='coerce').mean())

        row = {
            'Feature selector': fs_u,
            'Classifier': clf_u,
            'Features': features,

            'Training TN': tr.get('TN'),
            'Training FP': tr.get('FP'),
            'Training FN': tr.get('FN'),
            'Training TP': tr.get('TP'),
            'Training Accuracy': tr.get('Accuracy'),
            'CV Accuracy': cv_acc,
            'Training Precision': tr.get('Precision'),
            'Training Recall': tr.get('Recall'),
            'Training F1-score': tr.get('F1-score'),
            'Training BA': tr.get('BA'),
            'Training G-mean': tr.get('G-mean'),
            'Training ROC-AUC': tr.get('ROC-AUC'),

            'Test TN': te.get('TN'),
            'Test FP': te.get('FP'),
            'Test FN': te.get('FN'),
            'Test TP': te.get('TP'),
            'Test Accuracy': te.get('Accuracy', te.get('test_accuracy')),
            'Test Precision': te.get('Precision'),
            'Test Recall': te.get('Recall'),
            'Test F1-score': te.get('F1-score'),
            'Test BA': te.get('BA'),
            'Test G-mean': te.get('G-mean'),
            'Test ROC-AUC': te.get('ROC-AUC', te.get('test_auc')),
        }

        if row['Training Accuracy'] is not None and row['Test Accuracy'] is not None:
            row['Generalisation gap'] = float(row['Training Accuracy']) - float(row['Test Accuracy'])
        else:
            row['Generalisation gap'] = None

        rows.append(row)

    if not rows:
        print(" No complete model results are available to generate Table S12")
        return None

    order_map = {
        (fs, clf): i
        for i, (fs, clf) in enumerate(
            [(fs, clf) for fs in fs_order for clf in classifier_order]
        )
    }
    rows.sort(
        key=lambda r: order_map.get(
            (r['Feature selector'], r['Classifier']),
            999
        )
    )

    for i, row in enumerate(rows, start=1):
        row['No.'] = i

    columns = [
        'No.', 'Feature selector', 'Classifier', 'Features',
        'Training TN', 'Training FP', 'Training FN', 'Training TP',
        'Training Accuracy', 'CV Accuracy', 'Training Precision', 'Training Recall',
        'Training F1-score', 'Training BA', 'Training G-mean', 'Training ROC-AUC',
        'Test TN', 'Test FP', 'Test FN', 'Test TP',
        'Test Accuracy', 'Test Precision', 'Test Recall', 'Test F1-score',
        'Test BA', 'Test G-mean', 'Test ROC-AUC', 'Generalisation gap'
    ]

    df_s12 = pd.DataFrame(rows)[columns]
    os.makedirs(output_dir, exist_ok=True)

    csv_path = os.path.join(output_dir, 'Table_S12_generated.csv')
    df_s12.to_csv(csv_path, index=False, encoding='utf-8-sig')

    xlsx_path = os.path.join(output_dir, 'Table_S12_generated.xlsx')
    with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
        df_s12.to_excel(writer, sheet_name='Sheet1', index=False, startrow=1)
        ws = writer.book['Sheet1']

        from openpyxl.styles import Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        title_cell = ws.cell(row=1, column=1)
        title_cell.value = (
            'Table S12 Performance metrics of all models under each feature selection strategy '
            'on the training and test sets'
        )
        title_cell.font = Font(name='Arial', bold=True, size=11)
        title_cell.alignment = Alignment(horizontal='left', vertical='center')

        for cell in ws[2]:
            cell.font = Font(name='Arial', bold=True, size=9)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        integer_cols = {
            'No.', 'Features',
            'Training TN', 'Training FP', 'Training FN', 'Training TP',
            'Test TN', 'Test FP', 'Test FN', 'Test TP'
        }
        metric_cols = set(columns) - integer_cols - {'Feature selector', 'Classifier'}

        header_to_col = {
            ws.cell(row=2, column=c).value: c
            for c in range(1, len(columns) + 1)
        }
        first_data_row = 3
        last_data_row = 2 + len(df_s12)

        for r in range(first_data_row, last_data_row + 1):
            for c in range(1, len(columns) + 1):
                cell = ws.cell(r, c)
                cell.font = Font(name='Arial', size=9)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                header = ws.cell(2, c).value
                if header in metric_cols and isinstance(cell.value, (int, float, np.integer, np.floating)):
                    cell.number_format = '0.0000'
                elif header in integer_cols:
                    cell.number_format = '0'

        test_metric_headers = [
            'Test Accuracy', 'Test Precision', 'Test Recall', 'Test F1-score',
            'Test BA', 'Test G-mean', 'Test ROC-AUC'
        ]
        for header in test_metric_headers:
            c = header_to_col[header]
            vals = []
            for r in range(first_data_row, last_data_row + 1):
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    vals.append(v)
            if vals:
                best = max(vals)
                for r in range(first_data_row, last_data_row + 1):
                    if ws.cell(r, c).value == best:
                        ws.cell(r, c).font = Font(name='Arial', size=9, bold=True)

        thin = Side(style='thin', color='808080')
        for r in range(2, last_data_row + 1):
            for c in range(1, len(columns) + 1):
                ws.cell(r, c).border = Border(top=thin, bottom=thin)

        widths = {1: 6, 2: 16, 3: 12, 4: 10}
        for c in range(5, len(columns) + 1):
            widths[c] = 14
        for c, width in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = width

        ws.freeze_panes = 'A3'

    print(f" Table S12 saved to: {xlsx_path}")
    print(f" Table S12 CSV saved to: {csv_path}")
    return xlsx_path

def get_feature_category(feat_name):

    if feat_name.startswith('Tmax') or feat_name.startswith('Tmin') or feat_name.startswith('AT'):
        return 'Temperature'
    elif feat_name.startswith('RH'):
        return 'RH'
    elif feat_name.startswith('R'):
        return 'R'
    else:
        return None

def main():
    
    
    

    base_config = {
        'train_file': os.path.join(PROJECT_DIR, 'results', '03_feature_selection', 'data', 'features_selected_train.csv'),
        'test_file': os.path.join(PROJECT_DIR, 'results', '03_feature_selection', 'data', 'features_selected_test.csv'),
        'n_folds': 5,
        'random_state': 42,
        'pca_variance_threshold': 0.95,
        'result_dir': 'unified_results',
        'perform_grid_search': True,
        'threshold': 0.50
    }

    combinations = [
        ('rf', 'rf'), ('rf', 'mlp'), ('rf', 'dt'), ('rf', 'svm'), ('rf', 'lgb'), ('rf', 'xgb'),
        ('rfe', 'rf'), ('rfe', 'mlp'), ('rfe', 'dt'), ('rfe', 'svm'), ('rfe', 'lgb'), ('rfe', 'xgb')
    ]

    spore_engineer_params = {
        'do_spore_engineering': True,
        'spore_col': 'Spore_Index',
        'interact_cols': None,
        'log_transform': True
    }

    temp_df = pd.read_csv(base_config['train_file'])
    if 'target' in temp_df.columns:
        label_col = 'target'
    elif 'Label' in temp_df.columns:
        label_col = 'Label'
    else:
        label_col = temp_df.columns[-1]

    numeric_feats = temp_df.select_dtypes(include=[np.number]).columns.tolist()
    if label_col in numeric_feats:
        numeric_feats.remove(label_col)
    spore_col = spore_engineer_params['spore_col']
    if spore_col in numeric_feats:
        numeric_feats.remove(spore_col)
    exclude_non_meteor = ['CITY', 'county', 'year']
    for col in exclude_non_meteor:
        if col in numeric_feats:
            numeric_feats.remove(col)
    all_meteor_features = numeric_feats
    

    y_orig = temp_df[label_col].values.astype(int)

    spearman_vals = {}
    mi_vals = {}
    for feat in all_meteor_features:
        x_vals = temp_df[feat].values.astype(float)
        mask = ~np.isnan(x_vals)
        if mask.sum() == 0:
            spearman_vals[feat] = 0.0
            mi_vals[feat] = 0.0
            continue
        corr, _ = spearmanr(x_vals[mask], y_orig[mask])
        spearman_vals[feat] = abs(corr) if not np.isnan(corr) else 0.0
        mi = mutual_info_classif(x_vals[mask].reshape(-1, 1), y_orig[mask], random_state=42)[0]
        mi_vals[feat] = mi if not np.isnan(mi) else 0.0

    def minmax_norm(vals):
        minv, maxv = min(vals), max(vals)
        if maxv - minv == 0:
            return [0.0] * len(vals)
        return [(v - minv) / (maxv - minv) for v in vals]

    norm_spearman = {f: v for f, v in zip(all_meteor_features, minmax_norm(list(spearman_vals.values())))}
    norm_mi = {f: v for f, v in zip(all_meteor_features, minmax_norm(list(mi_vals.values())))}

    weight_s, weight_m = 0.5, 0.5
    combined_score = {f: weight_s * norm_spearman[f] + weight_m * norm_mi[f] for f in all_meteor_features}

    category_best = {}
    for feat in all_meteor_features:
        cat = get_feature_category(feat)
        if cat is None:
            continue
        if cat not in category_best or combined_score[feat] > combined_score[category_best[cat]]:
            category_best[cat] = feat

    selected_interact_cols = list(category_best.values())
    
    

    X_train, y_train, feature_names, _ = load_data_and_adjust_labels(
        base_config['train_file'],
        do_spore_engineering=spore_engineer_params['do_spore_engineering'],
        spore_col=spore_col,
        interact_cols=selected_interact_cols,
        is_train=True
    )
    if X_train is None:
        print(" Failed to load training data; terminating the script.")
        return

    n_total = len(feature_names)
    

    if os.path.exists(base_config['test_file']):
        X_test, y_test, _, _ = load_data_and_adjust_labels(
            base_config['test_file'],
            expected_feature_names=feature_names,
            do_spore_engineering=True,
            spore_col=spore_col,
            interact_cols=selected_interact_cols,
            is_train=False
        )
        if X_test is not None:
            pass
        else:
            print(" Failed to load the test set; test evaluation will be skipped.")
            X_test, y_test = None, None
    else:
        print(" Test-set file not found; skipping test evaluation.")
        X_test, y_test = None, None

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    total_dir = os.path.join(PROJECT_DIR, "results", f"model_comparison_{timestamp}")
    os.makedirs(total_dir, exist_ok=True)
    

    by_model_dir = os.path.join(total_dir, "by_model")
    by_metric_dir = os.path.join(total_dir, "by_metric")
    for d in [by_model_dir, by_metric_dir]:
        os.makedirs(d, exist_ok=True)
    for sub in ["roc_curves", "confusion_matrices", "feature_importance", "pca_data",
                "smote_distributions", "learning_curves", "permutation_importance", "pr_curves"]:
        os.makedirs(os.path.join(by_metric_dir, sub), exist_ok=True)

    result_paths = []
    all_raw_importances = {}
    all_shap_importances = {}

    for fs, clf in combinations:
        
        
        

        config = base_config.copy()
        config['feature_selector'] = fs
        config['classifier'] = clf

        model_key = f"{fs}+{clf}"
        model_subdir = os.path.join(by_model_dir, model_key)
        os.makedirs(model_subdir, exist_ok=True)

        try:

            fold_accs, mean_acc, core_features, all_results, fold_aucs = kfold_cross_validation(
                X_train, y_train, feature_names, config, model_subdir, by_metric_dir, threshold=config['threshold'])
            

            final_model, pca_analyzer, selector, core_indices, train_acc, raw_perm_imp, selected_features = train_final_model(
                X_train, y_train, core_features, feature_names, config, model_subdir, by_metric_dir
            )
            
            

            selected_indices = [core_features.index(feat) for feat in selected_features]
            X_train_final_selected = X_train[:, selected_indices]

            raw_perm_imp, shap_values, shap_feature_names = analyze_factor_importance(
                final_model,
                pca_analyzer,
                X_train_final_selected,
                y_train,
                X_train_final_selected,
                y_train,
                selected_features,
                feature_names,
                clf,
                model_subdir,
                by_metric_dir,
                model_key,
                n_permutations=30
            )

            all_raw_importances[model_key] = raw_perm_imp
            all_shap_importances[model_key] = {
                'shap_values': shap_values,
                'feature_names': shap_feature_names
            } if shap_values is not None else None

            test_acc, test_auc = None, None
            if X_test is not None and y_test is not None:
                test_acc, test_auc = evaluate_final_model_on_test_set(
                    final_model, pca_analyzer, selector, core_indices,
                    X_test, y_test, model_subdir, clf, by_metric_dir, model_key,
                    threshold=config['threshold'])
                if test_acc is not None:
                    pass

            generate_final_report(model_subdir, config, n_total, len(core_features),
                                  mean_acc, train_acc, test_acc, test_auc, clf)
            create_performance_summary_plot(fold_accs, fold_aucs, model_subdir)
            save_training_config(config, model_subdir)

            result_paths.append(model_subdir)
            

        except Exception as e:
            print(f" Model combination {fs.upper()}+{clf.upper()} failed: {str(e)}")
            traceback.print_exc()
            with open(os.path.join(model_subdir, "error_log.json"), 'w', encoding='utf-8') as f:
                json.dump(convert_types({'error': str(e), 'traceback': traceback.format_exc()}), f, indent=4)
            result_paths.append(model_subdir)
            all_raw_importances[model_key] = None
            all_shap_importances[model_key] = None

    if result_paths:
        generate_all_models_comparison(result_paths, base_config, by_metric_dir, all_raw_importances,
                                       all_shap_importances)

        generate_table_s12(result_paths, by_metric_dir)

    
    print(" All 12 model combinations have completed.")
    

if __name__ == "__main__":
    np.random.seed(42)
    main()
